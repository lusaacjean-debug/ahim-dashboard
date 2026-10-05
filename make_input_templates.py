#!/usr/bin/env python3
"""Create AHIM_Inputs.xlsx: one sheet per manual input, with dropdowns and a guide.
Usage: python scripts/make_input_templates.py [AHIM_Inputs.xlsx]"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter as L
import ahim_inputs as AI
OUT = sys.argv[1] if len(sys.argv) > 1 else 'AHIM_Inputs.xlsx'
NAVY = '203646'; F = lambda **k: Font(name='Arial', size=k.pop('size', 10), **k)
wb = Workbook(); g = wb.active; g.title = 'Guide'; g.sheet_view.showGridLines = False
g['A1'] = 'AHIM INPUTS: monthly data for the Outcomes, Finance, Work management, Field and Governance pages'; g['A1'].font = F(bold=True, size=14, color=NAVY)
g['A2'] = 'Fill the sheets below, save this file as AHIM_Inputs_<month>.xlsx in the kit "inputs" folder and run RUN_MONTHLY_UPDATE.bat. Rows are merged by key: re-sending a row with the same key updates it, so you can keep one file and add to it every month.'
g['A2'].font = F(color='5D6E79'); g['A2'].alignment = Alignment(wrap_text=True); g.row_dimensions[2].height = 30
hdr = ['Sheet', 'What it feeds', 'Who fills it / source', 'Frequency', 'Key (one row per)']
feeds = {'Work_Orders': 'Work management (automatic from the Pronto export: use only for flags Ready / Needs shutdown, or WOs not in the export)', 'Production': 'Outcomes: availability, downtime, lost production; Finance: cost per unit',
         'Maint_Costs': 'Finance: cost vs budget, cost / RAV', 'Labour': 'Work management: planned work %, emergency %, overtime, backlog crew-weeks', 'RCA': 'Actions & RCA, bad actors, KPI tree',
         'Actions': 'Actions & RCA, overdue actions KPI', 'KPI_Tree': 'KPI tree and the Overview tiles (strategic tier): edit owners and targets only', 'Routes': 'Field page: routes due, printable route sheets'}
for j, h in enumerate(hdr, 1):
    c = g.cell(4, j, h); c.font = F(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', start_color=NAVY)
for i, (name, (title, desc, cols, key, owner, freq)) in enumerate(AI.SCHEMAS.items(), 5):
    for j, v in enumerate([name, feeds[name], owner, freq, ' + '.join(key)], 1):
        c = g.cell(i, j, v); c.font = F(); c.alignment = Alignment(wrap_text=True, vertical='top')
for col, w in zip('ABCDE', [16, 70, 32, 16, 24]): g.column_dimensions[col].width = w
r = 6 + len(AI.SCHEMAS)
for t in ['RULES', '1. Do not rename sheets or column headers (row 4). Add rows below the header.', '2. Dates as real Excel dates. Month = first day of the month (e.g. 01-Sep-2026).',
          '3. Pronto asset no. must exist in the AHIM Asset_Register. Circuit and Area names must match the Asset_Register areas.',
          '4. Dropdown columns accept only the listed values. Hours and USD are plain numbers.', '5. To correct a row, re-enter it with the same key: the newest value wins.']:
    c = g.cell(r, 1, t); c.font = F(bold=(t == 'RULES')); r += 1
for name, (title, desc, cols, key, owner, freq) in AI.SCHEMAS.items():
    ws = wb.create_sheet(name)
    ws['A1'] = 'AHIM input: ' + title; ws['A1'].font = F(bold=True, size=13, color=NAVY)
    ws['A2'] = f'{desc} Key: {" + ".join(key)}. Source: {owner}. Frequency: {freq}.'; ws['A2'].font = F(size=9, color='5D6E79')
    for j, (c, w, kind, lst) in enumerate(cols, 1):
        h = ws.cell(4, j, c); h.font = F(bold=True, color='FFFFFF', size=9); h.fill = PatternFill('solid', start_color=NAVY); h.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[L(j)].width = w
        rng = f'{L(j)}5:{L(j)}5000'
        if kind == 'list':
            d = DataValidation(type='list', formula1='"' + ','.join(AI.LISTS[lst]) + '"', allow_blank=True, showErrorMessage=True, error='Choose from the list'); ws.add_data_validation(d); d.add(rng)
        elif kind == 'date':
            d = DataValidation(type='date', operator='greaterThan', formula1='36526', allow_blank=True, showErrorMessage=True, error='Enter a date'); ws.add_data_validation(d); d.add(rng)
            for rr in range(5, 205): ws[f'{L(j)}{rr}'].number_format = 'mmm-yy' if c == 'Month' else 'dd-mmm-yy'
        elif kind == 'num':
            d = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True, showErrorMessage=True, error='Enter a number'); ws.add_data_validation(d); d.add(rng)
    ws.row_dimensions[4].height = 34; ws.freeze_panes = 'A5'
    if name == 'KPI_Tree':
        for i, row in enumerate(AI.KPI_TREE, 5):
            for j, v in enumerate(row, 1): ws.cell(i, j, v).font = F()
wb.save(OUT); print('Wrote', OUT)
