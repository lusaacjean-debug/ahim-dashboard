#!/usr/bin/env python3
"""Validate the AHIM workbook before it is published.

Usage:  python scripts/validate_data.py [data/AHIM_Data.xlsx]
Exit code 1 if errors are found (GitHub Actions then blocks the deploy).
Warnings are reported but do not block.
"""
import sys, datetime as dt
from collections import Counter
from openpyxl import load_workbook

PATH = sys.argv[1] if len(sys.argv) > 1 else 'data/AHIM_Data.xlsx'
HDR = 4
errors, warnings = [], []
E = lambda m: errors.append(m)
W = lambda m: warnings.append(m)

wb = load_workbook(PATH, data_only=False)

def rows(name, key):
    if name not in wb.sheetnames:
        return None, {}
    ws = wb[name]
    hdr = {str(c.value).strip(): i for i, c in enumerate(ws[HDR]) if c.value}
    if key not in hdr:
        E(f'{name}: column "{key}" not found in row {HDR}')
        return [], hdr
    out = []
    for r in ws.iter_rows(min_row=HDR + 1, values_only=True):
        v = r[hdr[key]]
        if v is None or str(v).strip() == '':
            continue
        out.append({h: r[i] for h, i in hdr.items()})
    return out, hdr

def lst(col):
    ws = wb['Lists']; out = []
    for r in range(2, ws.max_row + 1):
        v = ws[f'{col}{r}'].value
        if v is None or str(v).strip() == '': break
        out.append(str(v).strip())
    return out

for s in ['Asset_Register', 'Records', 'Lists']:
    if s not in wb.sheetnames: E(f'Required sheet "{s}" is missing')
if errors:
    print('\n'.join('ERROR   ' + e for e in errors)); sys.exit(1)

AREAS, CLASSES, SOURCES, TECHS, STAGES, STATUS, DIRS = (lst(c) for c in 'ABCDEFG')

# ---- settings
ws = wb['Lists']
weights = [ws[f'M{r}'].value for r in range(2, 7)]
if any(not isinstance(x, (int, float)) for x in weights) or abs(sum(weights) - 1) > 0.001:
    E(f'Lists: ACI weights must be numbers totalling 1.00 (found {weights})')

# ---- register
reg, rh = rows('Asset_Register', 'Pronto asset no.')
ids = [str(a['Pronto asset no.']).strip() for a in reg]
for k, n in Counter(ids).items():
    if n > 1: E(f'Asset_Register: Pronto asset no. {k} appears {n} times (must be unique)')
FACT = ['Safety & env. (1-5)', 'Production (1-5)', 'Redundancy (1-5)', 'Repair cost / MTTR (1-5)', 'Failure history (1-5)']
for a in reg:
    k = a['Pronto asset no.']
    if a.get('Area') not in AREAS: E(f'Asset {k}: area "{a.get("Area")}" not in Lists')
    if a.get('Asset class') and a['Asset class'] not in CLASSES: W(f'Asset {k}: asset class "{a["Asset class"]}" not in Lists')
    f = [a.get(x) for x in FACT]
    if any(not isinstance(x, (int, float)) or not 1 <= x <= 5 for x in f): E(f'Asset {k}: criticality factors must all be 1-5 (found {f})')
    if not any(str(a.get(t) or '').upper() == 'Y' for t in TECHS): W(f'Asset {k}: no applicable technique marked Y')
    if str(a.get('Statutory cert. required') or '').upper() == 'Y' and not a.get('Certificate expiry'):
        W(f'Asset {k}: statutory certificate required but no expiry date')
idset = set(ids)

# ---- records
recs, _ = rows('Records', 'Pronto asset no.')
today = dt.datetime.now()
for i, r in enumerate(recs, start=HDR + 1):
    k, where = str(r['Pronto asset no.']).strip(), f'Records row ~{i}'
    if k not in idset: E(f'{where}: asset {k} is not in the Asset_Register')
    d = r.get('Date')
    if not isinstance(d, (dt.datetime, dt.date)): E(f'{where}: Date missing or not a date'); continue
    if d > today + dt.timedelta(days=1): E(f'{where}: Date {d:%d-%b-%Y} is in the future')
    if r.get('Source (input)') not in SOURCES: E(f'{where}: Source "{r.get("Source (input)")}" not in Lists')
    if r.get('Technique') not in TECHS: E(f'{where}: Technique "{r.get("Technique")}" not in Lists')
    stage = r.get('Record stage')
    if stage and stage not in STAGES: E(f'{where}: Record stage "{stage}" not in Lists')
    insp = r.get('Inspector status')
    if insp and insp not in STATUS: E(f'{where}: Inspector status "{insp}" must be OK, Alert or Danger')
    nums = [r.get('Value'), r.get('Alert limit'), r.get('Danger limit')]
    measured = all(isinstance(x, (int, float)) for x in nums)
    if not measured and not insp: E(f'{where}: no status possible: enter Value + both limits, or an Inspector status')
    if measured:
        dr = r.get('Limit direction')
        if dr not in DIRS: E(f'{where}: Limit direction required for measured values')
        elif dr == 'Higher is worse' and nums[1] > nums[2]: E(f'{where}: Alert limit must be below Danger limit (Higher is worse)')
        elif dr == 'Lower is worse' and nums[1] < nums[2]: E(f'{where}: Alert limit must be above Danger limit (Lower is worse)')
        st = insp or (('Danger' if nums[0] <= nums[2] else 'Alert' if nums[0] <= nums[1] else 'OK') if dr == 'Lower is worse' else ('Danger' if nums[0] >= nums[2] else 'Alert' if nums[0] >= nums[1] else 'OK'))
    else:
        st = insp
    if st in ('Alert', 'Danger') and stage != 'No action':
        if not r.get('Finding'): W(f'{where}: {st} without a Finding')
        if not r.get('Recommendation'): W(f'{where}: {st} without a Recommendation')
        if not stage: W(f'{where}: {st} without a Record stage')
    if stage == 'Closed':
        if not r.get('Date closed'): E(f'{where}: stage Closed but no Date closed')
        if not r.get('Verified by'): W(f'{where}: closed without "Verified by"')
    if r.get('Date closed') and isinstance(r.get('Date closed'), (dt.datetime, dt.date)) and r['Date closed'] < d:
        E(f'{where}: Date closed is before the record date')

# ---- professional fields (optional; checked when filled)
CLASSES_LIB = set()
if 'Strategy_Library' in wb.sheetnames:
    ws_ = wb['Strategy_Library']
    CLASSES_LIB = {str(r[0]).strip() for r in ws_.iter_rows(min_row=5, values_only=True) if r and r[0]}
for a in reg:
    k = a['Pronto asset no.']; sc = a.get('Strategy class'); c = a.get('Consequence (1-5)')
    if sc and CLASSES_LIB and str(sc).strip() not in CLASSES_LIB: E(f'Asset {k}: strategy class "{sc}" not in Strategy_Library')
    if c not in (None, '') and (not isinstance(c, (int, float)) or not 1 <= c <= 5): E(f'Asset {k}: Consequence must be 1-5 (found {c})')
    if not sc: W(f'Asset {k}: no strategy class (strategy compliance cannot be measured)')
for i, r in enumerate(recs, start=HDR + 1):
    w_ = r.get('Execution window'); l_ = r.get('Likelihood override (1-5)')
    if w_ and w_ not in ('Online', 'Unit isolation', 'Plant shutdown'): E(f'Records row ~{i}: Execution window "{w_}" must be Online, Unit isolation or Plant shutdown')
    if l_ not in (None, '') and (not isinstance(l_, (int, float)) or not 1 <= l_ <= 5): E(f'Records row ~{i}: Likelihood override must be 1-5')

# ---- optional sheets
for name, key in [('Events', 'Pronto asset no.'), ('CM_Value', 'Pronto asset no.'), ('Prestart', 'Pronto asset no.')]:
    rs, _ = rows(name, key)
    if rs is None: W(f'Optional sheet {name} not present'); continue
    for r in rs:
        if str(r[key]).strip() not in idset: E(f'{name}: asset {r[key]} is not in the Asset_Register')
ev, _ = rows('Events', 'Pronto asset no.')
for r in ev or []:
    if r.get('Event type') not in ('Failure', 'Repair', 'PM', 'Shutdown'): E(f'Events: type "{r.get("Event type")}" must be Failure, Repair, PM or Shutdown')
ps, _ = rows('Prestart', 'Pronto asset no.')
for r in ps or []:
    if (r.get('Prestarts completed') or 0) > (r.get('Shifts operated') or 0): E(f'Prestart: completed > shifts for asset {r["Pronto asset no."]}')

print(f'AHIM data check: {PATH}')
print(f'  {len(reg)} assets, {len(recs)} records, {len(errors)} errors, {len(warnings)} warnings')
for e in errors: print('ERROR   ' + e)
for w in warnings[:200]: print('WARNING ' + w)
sys.exit(1 if errors else 0)
