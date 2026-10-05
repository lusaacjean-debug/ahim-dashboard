"""Writes the AHIM reference sheets and the professional data fields into a workbook (shared by the builders)."""
import os, re
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
import ahim_reference as R

NAVY = '203646'; thin = Side(style='thin', color='C9D3D9'); B = Border(left=thin, right=thin, top=thin, bottom=thin)
def _f(**k): k.setdefault('name', 'Arial'); k.setdefault('size', 10); return Font(**k)

REG_EXTRA = [('Strategy class', 26), ('Consequence (1-5)', 11), ('Service', 20), ('Bottleneck (Y/N)', 10), ('Replacement value (USD)', 13), ('Consequence basis', 46), ('Install year', 9), ('Design life (years)', 9)]
REC_EXTRA = [('Failure mode (ISO 14224)', 11), ('Failure mechanism (ISO 14224)', 12), ('Damage mechanism (API 571)', 26), ('Execution window', 14), ('Likelihood override (1-5)', 11)]
WINDOWS = ['Online', 'Unit isolation', 'Plant shutdown']

def area_code(a):
    m = re.match(r'^(\d{3})', str(a[4] or '')) or re.match(r'^(\d{3})', str(a[0] or ''))
    return int(m.group(1)) if m else None

def default_extras(a):
    code = area_code(a)
    f = a[9] if isinstance(a[9], (list, tuple)) and len(a[9]) >= 2 else [3, 3]
    strat = R.strategy_class(a[5], code, a[2], a[0])
    svc = R.service_of(code, a[5])
    base = max(f[0], f[1]); fl, why = R.consequence_floor(svc, strat)
    cons = max(base, fl)
    basis = (f'Default: service floor {fl} ({why})' if fl > base else f'Default: criticality factors (safety and environment {f[0]}, production {f[1]})')
    return [strat, cons, svc, '', '', basis, '', '']

def merge_extras(a, old):
    """Defaults, overridden by the previous workbook. Consequence is kept only when its basis was set by a person
    (any basis not starting with 'Default'); defaults are recalculated on every import."""
    d = default_extras(a); names = [h for h, _ in REG_EXTRA]; out = list(d)
    for k, h in enumerate(names):
        v = (old or {}).get(h)
        if v in (None, ''): continue
        if h in ('Consequence (1-5)', 'Consequence basis'): continue
        out[k] = v
    ob = (old or {}).get('Consequence basis'); oc = (old or {}).get('Consequence (1-5)')
    if ob and not str(ob).startswith('Default') and oc not in (None, ''):
        out[names.index('Consequence (1-5)')] = oc; out[names.index('Consequence basis')] = ob
    return out

def old_sheet_rows(path, name, ncol, start=5, headers=None):
    if not path or not os.path.exists(path): return None
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=True)
        if name not in wb.sheetnames: return None
        if headers:   # keep the user's version only if it has the current layout
            h = [str(c).strip() if c else '' for c in next(wb[name].iter_rows(min_row=4, max_row=4, max_col=len(headers), values_only=True))]
            if h != list(headers): return None
        rows = [list(r) for r in wb[name].iter_rows(min_row=start, max_col=ncol, values_only=True) if r and r[0] not in (None, '')]
        return rows or None
    except Exception:
        return None

def _sheet(wb, name, title, sub, cols, rows, tab='9CB6C6', wrap=()):
    ws = wb.create_sheet(name)
    ws['A1'] = title; ws['A1'].font = _f(bold=True, size=14, color=NAVY)
    ws['A2'] = sub; ws['A2'].font = _f(size=9, color='5D6E79')
    for j, (h, w) in enumerate(cols, 1):
        c = ws.cell(4, j, h); c.font = _f(bold=True, color='FFFFFF', size=9); c.fill = PatternFill('solid', start_color=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical='center'); c.border = B
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[4].height = 32
    for i, r in enumerate(rows, 5):
        for j, v in enumerate(r, 1):
            c = ws.cell(i, j, v); c.font = _f(); c.border = B
            c.alignment = Alignment(wrap_text=(j in wrap), vertical='top')
    ws.freeze_panes = 'A5'; ws.sheet_properties.tabColor = tab
    return ws

def _lib(old_path):
    """User edits to the library are kept; classes or tasks added to the reference since are appended."""
    old = old_sheet_rows(old_path, 'Strategy_Library', 11)
    if not old: return [list(x) for x in R.STRATEGY]
    have = {(str(r[0]), str(r[3]), str(r[4])) for r in old}
    return old + [list(x) for x in R.STRATEGY if (x[0], x[3], x[4]) not in have]

def write_reference(wb, old_path=None):
    """Reference sheets: kept from the previous workbook when present (user edits win), else defaults."""
    keep = lambda name, n, default, headers=None: old_sheet_rows(old_path, name, n, headers=headers) or default
    HC = [('Level', 7), ('Name', 13), ('Safety', 22), ('Health and radiation', 40), ('Environment', 34), ('Production', 28), ('Financial', 14), ('Regulatory and legal', 36), ('Community and reputation', 24)]
    _sheet(wb, 'Risk_Consequence', 'Risk matrix: consequence levels (uranium process and sulphuric acid plant)',
           'Worst credible consequence across seven categories; the highest category sets the level. Dose criteria per ICRP 103 / IAEA GSR Part 3. Financial bands to be calibrated with site finance.',
           HC, keep('Risk_Consequence', 9, [list(x) for x in R.RISK_CONSEQUENCE], [h for h, _ in HC]), wrap=(3, 4, 5, 6, 7, 8, 9))
    HL = [('Level', 7), ('Name', 15), ('Description', 36), ('Frequency guide', 26), ('AHIM mapping', 54)]
    _sheet(wb, 'Risk_Likelihood', 'Risk matrix: likelihood levels', 'Frequency guide for risk assessments; AHIM sets likelihood from condition (an open finding can override it: Records, Likelihood override).',
           HL, keep('Risk_Likelihood', 5, [list(x) for x in R.RISK_LIKELIHOOD], [h for h, _ in HL]), wrap=(3, 4, 5))
    HR = [('Rating', 11), ('Minimum score', 10), ('Required response', 70), ('Authority to accept the risk', 26)]
    _sheet(wb, 'Risk_Rating', 'Risk matrix: rating, response and acceptance authority', 'Risk score = consequence x likelihood (1 to 25). A risk may only be accepted, unmitigated, by the authority shown or above.',
           HR, keep('Risk_Rating', 4, [list(x) for x in R.RISK_RATING], [h for h, _ in HR]), wrap=(3,))
    FL = [('Service', 30), ('Minimum consequence', 12), ('Reason', 60), ('Applies to', 60)]
    _sheet(wb, 'Consequence_Floors', 'Consequence floors by service', 'Minimum consequence for equipment that holds or moves the fluid (loss of containment). Every asset handling dry uranium product takes the product floor. Read-only reference: change in scripts/ahim_reference.py.',
           FL, [[k, v[0], v[1], 'All assets in the area' if k == 'Uranium product (yellowcake)' else 'Pumps, tanks, vessels, piping, heat exchangers, sumps, thickeners, acid towers and coolers'] for k, v in R.SERVICE_FLOOR.items()], wrap=(3, 4))
    codes = [['ISO 14224 failure mode', c, d, ''] for c, d in R.ISO_FAILURE_MODES] + [['ISO 14224 failure mechanism', c, d, ''] for c, d in R.ISO_MECHANISMS] + \
            [['Site failure mode mapping', k, v[0] + ' / ' + v[1], 'CM workbook Failure Mode -> ISO 14224'] for k, v in R.SITE_FM_MAP.items()]
    _sheet(wb, 'Failure_Codes', 'ISO 14224 failure codes', 'Failure modes (what the failure looks like) and mechanisms (how it happens), ISO 14224:2016 Annex B, plus the mapping of site failure-mode text.',
           [('Type', 26), ('Code', 34), ('Description', 40), ('Note', 36)], keep('Failure_Codes', 4, codes))
    _sheet(wb, 'Damage_Mechanisms', 'Damage mechanisms (static equipment)', 'API 571 damage mechanisms and site mechanisms relevant to the uranium process plant and the sulphuric acid plant.',
           [('Mechanism', 34), ('Source', 20), ('Where it applies', 90)], keep('Damage_Mechanisms', 3, [list(x) for x in R.DAMAGE_MECHANISMS]), wrap=(3,))
    _sheet(wb, 'Strategy_Library', 'Equipment strategy library', 'Failure mode -> technique -> task -> interval by criticality class (A/B/C, days; blank = not required). Starting intervals: refine with P-F interval, history and OEM (ISO 17359). Tracked = N: done by operators, not measured in AHIM.',
           [('Strategy class', 30), ('Failure mode (ISO 14224)', 10), ('Mechanism / cause', 36), ('Technique', 20), ('Task', 46), ('Interval A (days)', 9), ('Interval B (days)', 9), ('Interval C (days)', 9),
            ('Alert / Danger criterion', 40), ('Reference', 26), ('Tracked in AHIM', 9)], _lib(old_path), tab='2E8A57', wrap=(3, 5, 9))
    _sheet(wb, 'Standards_Map', 'Standards compliance map', 'How AHIM implements each standard and what remains to be done.',
           [('Standard', 24), ('Clause', 10), ('Requirement', 46), ('AHIM implementation', 56), ('Status', 34)], keep('Standards_Map', 5, [list(x) for x in R.STANDARDS_MAP]), wrap=(3, 4, 5))

def add_register_validation(ws, first_col, nrows):
    from openpyxl.utils import get_column_letter as L
    cls = sorted({x[0] for x in R.STRATEGY})
    d = DataValidation(type='list', formula1='=Strategy_Library!$A$5:$A$400', allow_blank=True)  # any class in the library
    ws.add_data_validation(d); d.add(f'{L(first_col)}5:{L(first_col)}{nrows}')
    d = DataValidation(type='whole', operator='between', formula1='1', formula2='5', allow_blank=True); ws.add_data_validation(d); d.add(f'{L(first_col+1)}5:{L(first_col+1)}{nrows}')
    d = DataValidation(type='list', formula1='"Y,N"', allow_blank=True); ws.add_data_validation(d); d.add(f'{L(first_col+3)}5:{L(first_col+3)}{nrows}')

def add_record_validation(ws, first_col, nrows):
    from openpyxl.utils import get_column_letter as L
    d = DataValidation(type='list', formula1='"' + ','.join(c for c, _ in R.ISO_FAILURE_MODES) + '"', allow_blank=True); ws.add_data_validation(d); d.add(f'{L(first_col)}5:{L(first_col)}{nrows}')
    d = DataValidation(type='list', formula1='"' + ','.join(WINDOWS) + '"', allow_blank=True); ws.add_data_validation(d); d.add(f'{L(first_col+3)}5:{L(first_col+3)}{nrows}')
    d = DataValidation(type='whole', operator='between', formula1='1', formula2='5', allow_blank=True); ws.add_data_validation(d); d.add(f'{L(first_col+4)}5:{L(first_col+4)}{nrows}')

def site_fm(text):
    return R.SITE_FM_MAP.get((text or '').strip(), ('', ''))

def mech_from_text(t):
    t = (t or '').lower()
    for k, v in [('active loss of containment', 'Gasket / bolting degradation'), ('gasket', 'Gasket / bolting degradation'), ('bolt', 'Gasket / bolting degradation'),
                 ('concrete', 'Concrete acid attack'), ('foundation', 'Foundation settlement'), ('anchor', 'Foundation settlement'), ('pitting', 'Localized (pitting) corrosion'),
                 ('retention', 'Atmospheric corrosion'), ('ponding', 'Atmospheric corrosion'), ('coating', 'Coating breakdown'), ('weld', 'Weld cracking / weld defects'),
                 ('lining', 'Lining failure'), ('corrosion', 'Atmospheric corrosion')]:
        if k in t: return v
    return ''


# ---------------- professional inputs: Work_Orders, Production, Maint_Costs, Labour, RCA, Actions, KPI_Tree, Routes ----------------
import datetime as _dt
import ahim_inputs as AI

def _norm(v):
    if isinstance(v, (_dt.datetime, _dt.date)): return v.strftime('%Y-%m-%d')
    return '' if v is None else str(v).strip()

def _key(row, cols, key):
    out = []
    for k in key:
        v = row[cols.index(k)]
        if k == 'Month' and v not in (None, ''):
            v = _norm(v)[:7] if not isinstance(v, str) or len(v) >= 7 else v
        out.append(_norm(v))
    return tuple(out)

def read_input_files(paths):
    """Rows from input workbooks (sheets named like the schemas, header in row 4, or row 1 if row 4 is not a header)."""
    from openpyxl import load_workbook
    got = {}
    for p in paths:
        wb = load_workbook(p, data_only=True, read_only=True)
        for name in AI.SCHEMAS:
            if name not in wb.sheetnames: continue
            cols = [c for c, *_ in AI.SCHEMAS[name][2]]
            allrows = list(wb[name].iter_rows(values_only=True))
            hr = next((i for i, r in enumerate(allrows[:6]) if r and str(r[0] or '').strip() == cols[0]), None)
            if hr is None: continue
            hdr = [str(h).strip() if h is not None else '' for h in allrows[hr]]
            for r in allrows[hr + 1:]:
                if not r or r[0] in (None, ''): continue
                got.setdefault(name, []).append([r[hdr.index(c)] if c in hdr and hdr.index(c) < len(r) else None for c in cols])
    return got

def write_input_sheets(wb, new_rows, old_path=None, input_files=(), log=None):
    """Old rows (previous workbook) <- new rows (generated, e.g. from the Pronto export) <- input files; upsert by key."""
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter as L
    extra = read_input_files(input_files) if input_files else {}
    for name, (title, desc, colspec, key, owner, freq) in AI.SCHEMAS.items():
        cols = [c for c, *_ in colspec]
        merged = {}
        for src, rows in (('previous', old_sheet_rows(old_path, name, len(cols), headers=cols) or []), ('generated', new_rows.get(name, [])), ('input', extra.get(name, []))):
            n = 0
            for r in rows:
                r = list(r) + [None] * (len(cols) - len(r))
                if r[0] in (None, ''): continue
                k = _key(r, cols, key)
                if src == 'generated' and k in merged and name == 'Work_Orders':
                    old = merged[k]   # keep manual Ready / Needs shutdown / finding ref flags on Pronto rows
                    for c in ('Ready (Y/N)', 'Needs shutdown (Y/N)', 'AHIM finding ref'):
                        i = cols.index(c)
                        if r[i] in (None, '') and old[i] not in (None, ''): r[i] = old[i]
                merged[k] = r; n += 1
            if log and src == 'input' and n: log(name, f'{n} rows merged from input workbook')
        rows = list(merged.values())
        di = [i for i, (_, _, kind, _) in enumerate(colspec) if kind == 'date']
        for r in rows:
            for i in di:
                if isinstance(r[i], str) and len(r[i]) >= 10:
                    try: r[i] = _dt.datetime.strptime(r[i][:10], '%Y-%m-%d')
                    except ValueError: pass
        ws = _sheet(wb, name, 'AHIM ' + title, f'{desc} Source: {owner}. Frequency: {freq}. Key: {" + ".join(key)}.',
                    [(c, w) for c, w, *_ in colspec], rows, tab='C9A227', wrap=tuple(i + 1 for i, (c, w, k, l) in enumerate(colspec) if w >= 30))
        for i, (c, w, kind, lst) in enumerate(colspec, 1):
            col = L(i)
            if kind == 'date':
                for rr in range(5, 5 + len(rows)): ws[f'{col}{rr}'].number_format = 'dd-mmm-yy' if c != 'Month' else 'mmm-yy'
            if kind == 'list':
                d = DataValidation(type='list', formula1='"' + ','.join(AI.LISTS[lst]) + '"', allow_blank=True); ws.add_data_validation(d); d.add(f'{col}5:{col}{5 + len(rows) + 2000}')
            if kind == 'num' and 'USD' in c:
                for rr in range(5, 5 + len(rows)): ws[f'{col}{rr}'].number_format = '#,##0'
