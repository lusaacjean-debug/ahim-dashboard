#!/usr/bin/env python3
"""Build or extend the AHIM UT Register from parsed UT reports.

Usage:  python scripts/build_ut_register.py parsed_reports.json UT_Register.xlsx [observations.json]

parsed_reports.json : list of reports {file, tag, name, date, green, red, grid{row:[values]}, cols[], dft[[..]], author}
                      (produced by a report parser; one entry per tank per inspection round)
observations.json   : {"aliases": {report_tag: ahim_asset_no}, "findings": [[tag, category, Alert|Danger, finding, recommendation]], "notes": {tag: note}}
If UT_Register.xlsx exists, readings from new rounds are added and existing UT_Findings rows (with their status, WO and closure) are kept.
Thickness findings are generated automatically for tanks with points in the red band.
"""
import sys, os, json, datetime as dt
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

src, out = sys.argv[1], sys.argv[2]
obs = json.load(open(sys.argv[3])) if len(sys.argv) > 3 else {}
alias, notes = obs.get('aliases', {}), obs.get('notes', {})
reports = json.load(open(src))

def band(v, g, r): return 'Red' if v <= r else 'Yellow' if v < g else 'Green'

# keep existing content
old_read, old_find, old_tanks = [], [], []
if os.path.exists(out):
    wb0 = load_workbook(out)
    old_read = [list(r) for r in wb0['UT_Readings'].iter_rows(min_row=2, values_only=True) if r[0]]
    old_find = [list(r) for r in wb0['UT_Findings'].iter_rows(min_row=2, values_only=True) if r[0]]
    old_tanks = [list(r) for r in wb0['UT_Tanks'].iter_rows(min_row=2, values_only=True) if r[0]]
seen = {(str(r[1]), str(r[2])[:10]) for r in old_read}

tanks, reads, finds = old_tanks[:], old_read[:], old_find[:]
fid = max([int(str(f[0]).split('-')[-1]) for f in old_find] + [0])
def nid():
    global fid; fid += 1; return f'UT-{fid:03d}'
existing_keys = {(str(f[2]), str(f[4])[:10], str(f[5]), str(f[7])[:60]) for f in old_find}

for o in reports:
    tag, date = o['tag'], o['date']
    if (tag, date) in seen: continue
    aid = alias.get(tag, tag)
    d = dt.datetime.strptime(date, '%Y-%m-%d')
    g, r = o['green'], o['red']
    vals = []
    for row, vs in sorted(o['grid'].items(), key=lambda x: int(x[0])):
        for c, v in enumerate(vs):
            reads.append([aid, tag, d, int(row), o['cols'][c] if c < len(o['cols']) else f'P{c+1}', v, g, r, band(v, g, r), o['file']])
            vals.append((v, int(row), o['cols'][c] if c < len(o['cols']) else f'P{c+1}'))
    nred = sum(v <= r for v, _, _ in vals); nyel = sum(r < v < g for v, _, _ in vals)
    mn = min(vals) if vals else (None, None, None)
    dft = min(min(x) for x in o['dft']) if o.get('dft') else None
    tanks.append([aid, tag, o['name'], d, 'Resultant Consulting Engineers', o.get('author', ''), o['file'], g, r,
                  f"{len(o['grid'])} x {len(o['cols'])}", len(vals), mn[0], f'row {mn[1]} {mn[2]}' if mn[1] else '',
                  round(nyel / len(vals) * 100, 1) if vals else None, round(nred / len(vals) * 100, 1) if vals else None, dft, notes.get(tag, '')])
    if nred:
        sev = 'Danger' if mn[0] < 0.7 * r else 'Alert'
        fnd = (f"{nred} of {len(vals)} UT points in the red band (≤ {r} mm), {nyel} yellow; minimum {mn[0]} mm at row {mn[1]} (from top), {mn[2]}.")
        rec = ("Calculate the API 653 minimum required thickness (t-min) per shell course and assess fitness for service of the red zones; "
               "set CMLs at the red points and re-measure within 12 months to establish the corrosion rate.")
        if sev == 'Danger': rec = 'Restrict service and assess immediately: ' + rec[0].lower() + rec[1:] + ' Plan repair (patch or plate replacement).'
        key = (aid, str(d)[:10], 'Thickness', fnd[:60])
        if key not in existing_keys: finds.append([nid(), aid, tag, d, d, 'Thickness', sev, fnd, rec, None, 'Open', None, None, o['file']])
    for (t, cat, sev, fnd, rec) in [f for f in obs.get('findings', []) if f[0] == tag]:
        key = (aid, str(d)[:10], cat, fnd[:60])
        if key not in existing_keys: finds.append([nid(), aid, tag, d, d, cat, sev, fnd, rec, None, 'Open', None, None, o['file']])

wb = Workbook()
H = Font(bold=True, color='FFFFFF'); F = PatternFill('solid', start_color='203646'); Y = PatternFill('solid', start_color='FFF8DC')
def sheet(name, cols, rows, widths, datecols=(), editable=()):
    ws = wb.create_sheet(name) if wb.sheetnames != ['Sheet'] else wb.active
    ws.title = name
    for j, c in enumerate(cols, 1):
        x = ws.cell(1, j, c); x.font = H; x.fill = F; x.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[chr(64 + j) if j <= 26 else 'A' + chr(64 + j - 26)].width = widths[j - 1]
    for i, r in enumerate(rows, 2):
        for j, v in enumerate(r, 1):
            x = ws.cell(i, j, v)
            if j in datecols: x.number_format = 'dd-mmm-yy'
            if j in editable: x.fill = Y
            if isinstance(v, str) and len(v) > 40: x.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    return ws
sheet('UT_Tanks', ['AHIM asset no.', 'Tag in report', 'Description', 'Inspection date', 'Contractor', 'Technician', 'Report file', 'Green ≥ (mm)', 'Red ≤ (mm)', 'Grid (rows x points)', 'Points', 'Minimum (mm)', 'Minimum location', '% Yellow', '% Red', 'Min external DFT (µm)', 'Notes'],
      tanks, [14, 13, 30, 11, 18, 18, 34, 9, 9, 10, 7, 9, 16, 8, 8, 10, 50], datecols=(4,), editable=(1,))
sheet('UT_Readings', ['AHIM asset no.', 'Tag in report', 'Date', 'Elevation row (1 = top)', 'Location', 'Thickness (mm)', 'Green ≥ (mm)', 'Red ≤ (mm)', 'Band', 'Report file'],
      reads, [14, 13, 11, 10, 9, 10, 9, 9, 8, 34], datecols=(3,))
ws = sheet('UT_Findings', ['Finding ID', 'AHIM asset no.', 'Tag in report', 'Inspection date', 'Date raised', 'Category', 'Severity', 'Finding', 'Recommendation', 'Pronto WO #', 'Status', 'Date closed', 'Verified by', 'Source report'],
      finds, [10, 14, 13, 11, 11, 12, 9, 60, 60, 11, 10, 11, 14, 30], datecols=(4, 5, 12), editable=(10, 11, 12, 13))
dv = DataValidation(type='list', formula1='"Open,WO raised,Scheduled,Closed"', allow_blank=True); ws.add_data_validation(dv); dv.add('K2:K2000')
dv2 = DataValidation(type='list', formula1='"Alert,Danger"', allow_blank=False); ws.add_data_validation(dv2); dv2.add('G2:G2000')
g = wb.create_sheet('Guide', 0)
for i, t in enumerate([
    'AHIM UT Register: shell thickness surveys and their findings',
    '',
    'UT_Tanks: one row per tank per survey. Column A (AHIM asset no.) links the tank to the AHIM register: correct it if the report tag is wrong.',
    'UT_Readings: every thickness point (row 1 = top of tank). Add the next survey round as new rows with its date: AHIM then calculates the corrosion rate and remaining life per elevation band.',
    'UT_Findings: recommendations. Update Pronto WO #, Status, Date closed and Verified by (yellow columns) as work progresses. Severity: Danger = act now, Alert = plan.',
    '',
    'Bands are the contractor classification (Green acceptable, Yellow minor loss, Red significant loss). Red is NOT the API 653 minimum required thickness (t-min). Until t-min is calculated per shell course, AHIM uses: Alert = red band (assessment needed), Danger = below 70% of the red limit.',
    'Import into AHIM: python scripts/import_cm_workbook.py CM_Monthly_Report.xlsx AHIM_Data.xlsx UT_Register.xlsx [Vessels_...xlsx]']):
    c = g.cell(i + 1, 1, t); c.font = Font(bold=(i == 0), size=14 if i == 0 else 10); c.alignment = Alignment(wrap_text=True)
g.column_dimensions['A'].width = 140
wb.save(out)
print(f'{out}: {len(tanks)} tank surveys, {len(reads)} readings, {len(finds)} findings')
