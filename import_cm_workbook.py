#!/usr/bin/env python3
"""Import the site CM Monthly Report workbook into the AHIM data workbook.

Usage:  python scripts/import_cm_workbook.py CM_Monthly_Report.xlsx data/AHIM_Data.xlsx [Vessels_Tanks_Inspection_Summary.xlsx ...]
Optional extra files: AHIM input workbooks (sheets Work_Orders, Production, Maint_Costs, Labour, RCA, Actions, KPI_Tree, Routes), vessel & tank statutory inspection summaries (sheets 'Vessel Register' and 'Action Register') and the AHIM UT Register (UT_Tanks, UT_Readings, UT_Findings).
Re-run every month on the updated CM workbook. Conversion rules and assumptions are written to the Import_Review sheet.
"""
import sys as _sys, os as _os
_sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
import ahim_reference as AR_REF, ahim_sheets as AS, ahim_inputs as AI
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L
import datetime as dt

# ======================= EXTRACT FROM THE CM MONTHLY REPORT WORKBOOK =======================
import sys, os, re, warnings
import pandas as pd
warnings.filterwarnings('ignore')
SRC = sys.argv[1] if len(sys.argv) > 1 else 'CM_Monthly_Report.xlsx'
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'data', 'AHIM_Data.xlsx')
REVIEW = []   # (category, item, detail)
def rv(cat, item, detail=''): REVIEW.append((cat, str(item), str(detail)))

AREA_NAMES = {110:'Feed preparation (crushing)',120:'Milling',210:'Pre-leach thickener',220:'Leaching',310:'Resin-in-pulp (RIP)',320:'Elution',
 330:'Nano / membrane plant',410:'Gypsum precipitation',420:'Uranium precipitation',500:'Product scrubber',510:'Product recovery and drying',
 520:'Product baghouse and stack',610:'Neutralisation and tailings',620:'Tailings return water',660:'Neutralisation',710:'Lime',720:'Flocculant',
 730:'Hydrogen peroxide',740:'Caustic',750:'Dilute acid',760:'Magnesium oxide',770:'Magnesium oxide make-up',810:'Process water',
 820:'Raw and fire water',830:'Raw / filtered water distribution',840:'Potable water'}
GROUPS = {1:'Crushing & milling',2:'Leaching',3:'RIP & elution',4:'Precipitation & product',5:'Precipitation & product',6:'Tailings & neutralisation',7:'Reagents',8:'Water services'}
def area_of(code):
    try: c = int(code)
    except Exception: return 'Other', 'Unknown area'
    return GROUPS.get(c // 100, 'Other'), f'{c} {AREA_NAMES.get(c, "Area " + str(c))}'
def code_from_tag(t):
    m = re.match(r'^(\d{3})', str(t)); return int(m.group(1)) if m else None
def S(v): return '' if v is None or (isinstance(v, float) and pd.isna(v)) or (v is pd.NaT) else str(v).strip()
def D(v):
    v = pd.to_datetime(v, errors='coerce'); return None if pd.isna(v) else v.to_pydatetime().replace(hour=0, minute=0, second=0, microsecond=0)

# ---- Data entry
de = pd.read_excel(SRC, 'Data entry'); de = de[de['Equipment ID'].notna()].copy()
de['tag'] = de['Equipment ID'].astype(str).str.strip()
de['date'] = pd.to_datetime(de['Inspection Date'], errors='coerce')
bad_dates = de['date'].isna().sum()
de = de[de['date'].notna()].sort_values(['tag', 'date']).reset_index(drop=True)
MAXD = de['date'].max()
# ---- Criticality master
cr = pd.read_excel(SRC, 'Criticality'); cr = cr[cr['TAG'].notna()]
MASTER = {str(t).strip(): (int(c) if not pd.isna(c) else None, S(n)) for t, c, n in zip(cr['TAG'], cr['Criticality'], cr['DESC'])}
# ---- Oil
oil = pd.read_excel(SRC, 'Oil Analysis', header=3); oil = oil[oil['Machine ID'].notna() & oil['Date Sampled'].notna()].copy()
oil['tag'] = oil['Machine ID'].astype(str).str.strip()
# ---- Vibration
vib = pd.read_excel(SRC, 'Vibration Analysis', header=3).iloc[:, :40]
vib = vib[vib['Machine ID'].astype(str).str.match(r'^\d{3}-') & vib['Date Measured'].notna()].copy()
vib['tag'] = vib['Machine ID'].astype(str).str.strip()

# ---- Build asset register
tags = sorted(set(de['tag']) | set(oil['tag']) | set(vib['tag']))
tags = list(tags)
BASE = {1:[5,5,4,4], 2:[3,4,3,2], 3:[2,2,2,1], 4:[1,1,1,1]}
yr_ago = MAXD - pd.Timedelta(days=365)
findings12 = de[(de['date'] > yr_ago) & de['Status'].isin(['Alert', 'Danger'])].groupby('tag').size()
def fh(n): return 1 if n == 0 else 2 if n == 1 else 3 if n <= 3 else 4 if n <= 5 else 5
assets = []; types = set(); conflicts = []
for t in tags:
    g = de[de['tag'] == t]
    last = lambda col: next((S(v) for v in reversed(g[col].tolist()) if S(v)), '')
    dec = last('Criticality'); dec_n = int(dec[1]) if re.match(r'^C[1-4]$', dec) else None
    m = MASTER.get(t)
    crit = m[0] if m and m[0] else dec_n
    if m and m[0] and dec_n and m[0] != dec_n: conflicts.append((t, m[0], dec))
    if crit is None:
        crit = 2; rv('Criticality not assigned', t, 'No criticality in the Criticality sheet or Data entry. Provisional C2 used: confirm.')
    name = (m[1] if m and m[1] else '') or last('Equipment Name')
    if not name:
        o = oil[oil['tag'] == t]; v = vib[vib['tag'] == t]
        name = S(o['Description'].iloc[0]) if len(o) else S(v['Equipment Name'].iloc[0]) if len(v) else t
    code = None
    if len(g): 
        try: code = int(g['Area'].dropna().iloc[-1])
        except Exception: code = None
    code = code or code_from_tag(t)
    grp, sysloc = area_of(code)
    etype = last('Equipment Type') or 'Unclassified'; types.add(etype)
    techs = 'VI' + (' VA' if t in set(vib['tag']) else '') + (' OA' if t in set(oil['tag']) else '')
    f = BASE[crit] + [fh(int(findings12.get(t, 0)))]
    trade = last('Trades')
    assets.append((t, t, name, grp, sysloc, etype, '', '', 'Duty', f, techs, 'N', '', trade, f'C{crit}'))
    if t not in set(de['tag']): rv('Asset only in oil/vibration data', t, 'Added to register from oil or vibration sheet; not in Data entry.')
for t, mc, dc in conflicts: rv('Criticality conflict', t, f'Criticality sheet = {mc}; latest Data entry = {dc}. Criticality sheet used.')
# ---- Optional: vessel & tank statutory inspection summaries (extra arguments after OUT)
VESSEL_FILES = [a for a in sys.argv[3:] if a.lower().endswith(('.xlsx', '.xlsm'))]
VREC = []
def _hdr(path, sheet, key):
    raw = pd.read_excel(path, sheet, header=None)
    h = raw.index[raw.apply(lambda r: key in [str(x).strip() for x in r.values], axis=1)][0]
    return pd.read_excel(path, sheet, header=h)
# ---- Optional: AHIM UT Register (sheets UT_Tanks, UT_Readings, UT_Findings)
def _is_ut(path):
    try:
        from openpyxl import load_workbook as _l
        return 'UT_Readings' in _l(path, read_only=True).sheetnames
    except Exception: return False
def _is_input(path):
    try:
        from openpyxl import load_workbook as _l
        return bool(set(_l(path, read_only=True).sheetnames) & set(AI.SCHEMAS))
    except Exception: return False
INPUT_FILES = [f for f in VESSEL_FILES if _is_input(f) and not _is_ut(f)]
VESSEL_FILES = [f for f in VESSEL_FILES if f not in INPUT_FILES]
UT_FILES = [f for f in VESSEL_FILES if _is_ut(f)]
VESSEL_FILES = [f for f in VESSEL_FILES if f not in UT_FILES]
CAT2TECH = {'Thickness': 'Ultrasonic thickness', 'Structural': 'Structural', 'Foundation': 'Structural'}
for uf in UT_FILES:
    utk = pd.read_excel(uf, 'UT_Tanks'); utk = utk[utk['AHIM asset no.'].notna()]
    urd = pd.read_excel(uf, 'UT_Readings'); urd = urd[urd['AHIM asset no.'].notna() & urd['Thickness (mm)'].notna()]
    ufd = pd.read_excel(uf, 'UT_Findings'); ufd = ufd[ufd['AHIM asset no.'].notna()]
    for _, r in utk.drop_duplicates('AHIM asset no.', keep='last').iterrows():
        t = S(r['AHIM asset no.'])
        if t not in tags:
            code = code_from_tag(t); grp, sysloc = area_of(code); m = MASTER.get(t)
            crit = m[0] if m and m[0] else None
            if crit is None:
                crit = 2; rv('Criticality not assigned', t, f"UT tank {S(r['Description'])}: provisional C2 used, confirm")
            nf = int((ufd['AHIM asset no.'].astype(str).str.strip() == t).sum())
            assets.append((t, t, S(r['Description']) or t, grp, sysloc, 'Tank', '', '', 'Duty', BASE[crit] + [fh(nf)], 'UT VI TV STR', 'N', '', 'Mechanical', f'C{crit}'))
            tags.append(t); types.add('Tank')
        else:
            for i, a in enumerate(assets):
                if a[0] == t: assets[i] = a[:10] + (' '.join(sorted(set(a[10].split()) | {'UT', 'TV', 'STR'})),) + a[11:]
        if S(r.get('Tag in report')) and S(r.get('Tag in report')) != t:
            rv('UT tag linked', S(r['Tag in report']), f'Report tag linked to AHIM asset {t}: confirm')
    # one record per tank, survey date and elevation band: band minimum against the contractor bands
    for (t, d, row), g in urd.groupby(['AHIM asset no.', 'Date', 'Elevation row (1 = top)']):
        g = g.sort_values('Thickness (mm)'); mn = g.iloc[0]
        nred = int((g['Band'] == 'Red').sum()); nyel = int((g['Band'] == 'Yellow').sum())
        VREC.append((D(d).strftime('%Y-%m-%d'), S(t), 'Ultrasonic report', 'Ultrasonic thickness', f'Shell band {int(row):02d} from top: minimum',
                     float(mn['Thickness (mm)']), 'mm', float(mn['Red ≤ (mm)']), round(0.7 * float(mn['Red ≤ (mm)']), 2), 'Lower is worse', '',
                     '', '', '', 'No action', 'UT contractor', '', '',
                     f"Minimum at {S(mn['Location'])}; {len(g)} points, {nred} red, {nyel} yellow, band mean {g['Thickness (mm)'].mean():.2f} mm; Alert = contractor red band, Danger = 70% of red (provisional until API 653 t-min is calculated)"))
    for _, r in ufd.iterrows():
        t = S(r['AHIM asset no.']); cat = S(r['Category']); st = S(r['Status']).lower()
        dc = D(r.get('Date closed')); wo = S(r.get('Pronto WO #'))
        stage = 'Closed' if (st == 'closed' or dc) else 'Scheduled' if st == 'scheduled' else 'WO raised' if (wo or st == 'wo raised') else 'Raised'
        VREC.append(((D(r['Date raised']) or D(r['Inspection date'])).strftime('%Y-%m-%d'), t, 'Ultrasonic report', CAT2TECH.get(cat, 'Tank & vessel'),
                     f'{cat}: survey finding', '', '', '', '', '', S(r['Severity']) or 'Alert', S(r['Finding']), S(r['Recommendation']), wo, stage,
                     'UT contractor', dc.strftime('%Y-%m-%d') if (stage == 'Closed' and dc) else '', S(r.get('Verified by')) or ('Not recorded' if stage == 'Closed' else ''),
                     f"{S(r['Finding ID'])}; {S(r['Source report'])}"))
    rv('Summary', 'UT register', f"{os.path.basename(uf)}: {utk['AHIM asset no.'].nunique()} tanks, {len(urd)} readings condensed to elevation bands, {len(ufd)} findings ({int((ufd['Status'].astype(str).str.lower()=='closed').sum())} closed)")

RATING = {1: 'OK', 2: 'OK', 3: 'Alert', 4: 'Danger', 5: 'Danger'}
for vf in VESSEL_FILES:
    try:
        vr = _hdr(vf, 'Vessel Register', 'Equipment Tag'); ar_ = _hdr(vf, 'Action Register', 'Action ID')
    except Exception as e:
        rv('Vessel inspection', os.path.basename(vf), f'Could not read ({e})'); continue
    vr = vr[vr['Equipment Tag'].notna() & vr['Inspection Date'].notna()]
    ar_ = ar_[ar_['Action ID'].astype(str).str.startswith('ACT')]
    nv = na_ = 0
    for _, r in vr.iterrows():
        t = S(r['Equipment Tag'])
        if not re.match(r'^\d{3}-', t):
            rv('Vessel not linked', t, f"{S(r['Vessel'])} inspected {D(r['Inspection Date']):%d-%b-%Y} has no valid tag: confirm the tag and re-import"); continue
        nv += 1
        nact = int((ar_['Equipment Tag'].astype(str).str.strip() == t).sum())
        nxt = D(r.get('Next Inspection Due'))
        if t not in tags:
            code = code_from_tag(t); grp, sysloc = area_of(code); m = MASTER.get(t)
            crit = m[0] if m and m[0] else None
            if crit is None:
                crit = 2; rv('Criticality not assigned', t, f"Vessel {S(r['Vessel'])}: provisional C2 used, confirm (acid/eluate service and leak consequence may justify C1)")
            nm = S(r['Vessel']); nm = f'{nm} (vessel/tank)' if nm and 'not' not in nm.lower() else 'Vessel / tank (name not stated)'
            assets.append((t, t, nm, grp, sysloc, 'Vessel / tank', '', '', 'Duty', BASE[crit] + [fh(nact)], 'VI TV STR ST_', 'Y',
                           nxt.strftime('%Y-%m-%d') if nxt else '', 'Mechanical', f'C{crit}'))
            tags.append(t); types.add('Vessel / tank')
        else:  # existing asset: add statutory due date and integrity techniques
            for i, a in enumerate(assets):
                if a[0] == t:
                    assets[i] = a[:10] + (' '.join(sorted(set(a[10].split()) | {'VI', 'TV', 'STR', 'ST_'})), 'Y', nxt.strftime('%Y-%m-%d') if nxt else a[12]) + a[13:]
        worst = r.get('Worst Rating'); worst = int(worst) if not pd.isna(worst) else 3
        VREC.append((D(r['Inspection Date']).strftime('%Y-%m-%d'), t, 'Statutory inspection', 'Tank & vessel', 'External visual inspection (overall)',
                     '', '', '', '', '', RATING.get(worst, 'Alert'),
                     f"Overall {S(r['Overall Condition'])} (worst component rating {worst}); {S(r['Fitness for Service (visual)'])}; {S(r['Checks NOT OK'])} checks NOT OK of {S(r['Total Checks'])}",
                     '', '', 'No action', S(r['Inspected By']), '', '',
                     f"Reviewed by {S(r['Reviewed By'])}; method {S(r['Method'])}; reported health {S(r['Health % (reported)'])}; next inspection due {nxt:%d-%b-%Y}" if nxt else ''))
    for _, r in ar_.iterrows():
        t = S(r['Equipment Tag'])
        if not re.match(r'^\d{3}-', t): continue
        na_ += 1
        sec = S(r['Section']).lower(); tech = 'Structural' if 'foundation' in sec or 'structur' in sec else 'Tank & vessel'
        insp = 'Danger' if S(r['Priority']) == 'P1' else 'Alert'
        stt = S(r['Status']); dc = D(r.get('Date Closed')); wo = S(r.get('Pronto WO #'))
        stage = 'Closed' if (stt.lower() == 'closed' or dc) else ('WO raised' if wo else 'Raised')
        nd = S(r.get('NDT / Verification'))
        VREC.append((D(r['Inspection Date']).strftime('%Y-%m-%d'), t, 'Statutory inspection', tech, S(r['Checklist Item']) or S(r['Section']),
                     '', '', '', '', '', insp, S(r['Observation (as recorded)']), S(r['Required Action']) + (f' Verification: {nd}.' if nd else ''),
                     wo, stage, S(r.get('Responsible')) or 'Vessel inspection team', dc.strftime('%Y-%m-%d') if (stage == 'Closed' and dc) else '',
                     'Not recorded' if stage == 'Closed' else '',
                     f"{S(r['Action ID'])}; report priority {S(r['Priority'])} (risk {S(r['Risk Score'])}), target {D(r['Target Date']):%d-%b-%Y}; {S(r['Defect Category'])}; {S(r['Damage Mechanism / Cause'])}; {S(r['Finding Basis'])}; {S(r['Discipline'])}" if D(r.get('Target Date')) else S(r['Action ID'])))
    rv('Summary', 'Vessel inspection file', f'{os.path.basename(vf)}: {nv} vessels, {na_} actions imported (source "Statutory inspection"). Next inspection due dates feed the statutory register.')

# ---- Keep manual edits from the previous AHIM workbook (register corrections, manual sheets)
OLD = {}
OLDX = {}
KEPT_ASSETS = 0
if os.path.exists(OUT):
    try:
        from openpyxl import load_workbook as _lw
        _ow = _lw(OUT, data_only=True, read_only=True)
        def _rows(name, ncol):
            if name not in _ow.sheetnames: return []
            out = []
            for r in _ow[name].iter_rows(min_row=5, max_col=ncol, values_only=True):
                if r[0] is None or str(r[0]).strip() == '': continue
                out.append(list(r))
            return out
        for nm, nc in [('CM_Value', 7), ('Programme_Cost', 5), ('Decisions', 5), ('Prestart', 4), ('Commentary', 5)]:
            OLD[nm] = _rows(nm, nc)
        _reg = pd.read_excel(OUT, 'Asset_Register', header=3)
        _reg = _reg[_reg['Pronto asset no.'].notna()]
        TK = {'Vibration':'VA','Lubrication / oil':'OA','Ultrasonics (airborne)':'US','Ultrasonic thickness':'UT','Infrared':'IR','Visual':'VI','Statutory':'ST_','Tank & vessel':'TV','Piping':'PI','Structural':'STR','Electrical':'EI','DG checks':'DG','Fleet inspection':'FI','Prestart':'PS'}
        oldreg = {str(r['Pronto asset no.']).strip(): r for _, r in _reg.iterrows()}
        for _k, _o in oldreg.items():
            OLDX[_k] = {h: (S(_o.get(h)) if not isinstance(_o.get(h), (int, float)) else (None if pd.isna(_o.get(h)) else _o.get(h))) for h, _ in AS.REG_EXTRA if h in _o}
        new = []
        for a in assets:
            o = oldreg.get(a[0])
            if o is None: new.append(a); continue
            g = lambda c, d: S(o.get(c)) or d
            fac = [o.get(c) for c in ['Safety & env. (1-5)', 'Production (1-5)', 'Redundancy (1-5)', 'Repair cost / MTTR (1-5)', 'Failure history (1-5)']]
            fac = [int(x) for x in fac] if all(isinstance(x, (int, float)) and not pd.isna(x) for x in fac) else a[9]
            tk = set(a[10].split()) | {TK[t] for t in TK if t in o and S(o.get(t)).upper() == 'Y'}
            exp = D(o.get('Certificate expiry')); exp = exp.strftime('%Y-%m-%d') if exp else ''
            if a[12] and (not exp or a[12] > exp): exp = a[12]   # newer inspection due date wins
            new.append((a[0], g('Site tag', a[1]), g('Description', a[2]), g('Area', a[3]), g('System / location', a[4]), g('Asset class', a[5]),
                        g('Manufacturer', ''), g('Model', ''), g('Operating mode', a[8]), fac, ' '.join(sorted(tk)), ('Y' if a[11] == 'Y' else g('Statutory cert. required', 'N')), exp,
                        g('Responsible', a[13]), a[14]))
            KEPT_ASSETS += 1
        assets = new
        types |= {a[5] for a in assets}
        rv('Summary', 'Manual edits kept', f'{KEPT_ASSETS} register rows kept your previous edits (area, class, criticality factors, statutory, responsible); manual sheets carried over: ' + ', '.join(f'{k} ({len(v)})' for k, v in OLD.items()))
    except Exception as e:
        rv('Summary', 'Manual edits kept', f'Could not read previous workbook ({e}); starting fresh')

AREA_LIST = sorted({a[3] for a in assets}) + [x for x in ['Acid plant', 'Power plant', 'Mobile fleet', 'Utilities', 'Other'] if x not in {a[3] for a in assets}]
TYPE_LIST = sorted(t for t in types if t) + [x for x in ['Tank', 'Pressure vessel', 'Piping', 'Structure', 'Boiler', 'Fired equipment', 'Electrical - transformer', 'Electrical - switchgear / MCC', 'Diesel generator', 'Light vehicle', 'Heavy equipment'] if x not in types]

# ---- Records
recs = []; n_off = n_nostat = n_untracked = n_est = n_rec = 0
ST = {'Ok':'OK', 'Alert':'Alert', 'Danger':'Danger'}
alld = de.groupby('tag')['date'].apply(list).to_dict()
for _, r in de.iterrows():
    st = S(r['Status'])
    if st == 'Offline': n_off += 1; continue
    if st not in ST: n_nostat += 1; continue
    insp = ST[st]; dte = r['date'].to_pydatetime(); rs = S(r['Record Status']); wo = S(r['WR/WO  Number']); wrs = S(r['WR/WO  Status'])
    rem = '; '.join(x for x in [('Failure mode: ' + S(r['Failure Mode'])) if S(r['Failure Mode']) else '',
        (f"Risk {S(r['Risk Score'])} ({S(r['Risk Level'])})") if S(r['Risk Score']) else '', ('Trade: ' + S(r['Trades'])) if S(r['Trades']) else '',
        ('WR/WO status: ' + wrs) if wrs else '', ('Equipment ' + S(r['Equipment status']).lower()) if S(r['Equipment status']) else ''] if x)
    closed = verified = ''; stage = 'No action'
    if insp != 'OK':
        if rs == 'Open': stage = 'WO raised' if (wo or wrs == 'Approved') else 'Raised'
        elif rs == 'Closed':
            stage = 'Closed'
            dc_ = D(r.get('Date Closed')) if 'Date Closed' in de.columns else None
            if dc_:   # recorded closure (CM workbook columns Date Closed / Verified By)
                closed = dc_.strftime('%Y-%m-%d'); verified = S(r.get('Verified By')) or 'Not recorded'; n_rec += 1
            else:
                nxt = [x for x in alld[r['tag']] if x > r['date']]
                closed = (nxt[0] if nxt else r['date']).strftime('%Y-%m-%d'); verified = 'Re-inspection'; n_est += 1
                rem += ('; ' if rem else '') + ('Close date = next inspection (estimated)' if nxt else 'Close date not recorded')
        else:
            n_untracked += 1
            rem += ('; ' if rem else '') + f'Record status "{rs or "blank"}" in source: kept as a reading, not tracked'
            rv('Alert/Danger not tracked', r['tag'], f"{dte:%d-%b-%Y} {insp}, record status '{rs or 'blank'}': {S(r['Findings/Observation'])[:120]}")
    recs.append((dte.strftime('%Y-%m-%d'), r['tag'], 'Vibration & visual route', 'Visual', 'General condition (sensitive inspection)',
                 '', '', '', '', '', insp, S(r['Findings/Observation']), S(r['Action Required']) if insp != 'OK' else '', wo, stage,
                 S(r['Inspector']), closed, verified, rem))
OILST = {'NORMAL':'OK', 'BORDERLINE':'Alert', 'URGENT':'Danger', 'CRITICAL':'Danger', 'ABNORMAL':'Alert'}
for _, r in oil.iterrows():
    insp = OILST.get(S(r['Severity']).upper(), 'Alert'); act = S(r['Recommended Action'])
    stage = 'No action' if insp == 'OK' or (insp == 'Alert' and re.match(r'^(routine|monitor)', act.lower())) else 'Raised'
    fe = r['Fe (Iron) ppm']; fe = float(fe) if not pd.isna(fe) else ''
    rem = f"Sample {S(r['Sample No.'])}; {S(r['Lubricant (Oil)'])}; PQ {S(r['PQ Index'])}; Si {S(r['Si'])}; water {S(r['Water %'])}%; visc40 {S(r['Viscosity @40C cSt'])} cSt; TAN {S(r['TAN (mgKOH/g)'])}"
    recs.append((D(r['Date Sampled']).strftime('%Y-%m-%d'), r['tag'], 'Oil analysis', 'Lubrication / oil', f"{S(r['Component']).title()}: iron (Fe)",
                 fe, 'ppm' if fe != '' else '', '', '', '', insp, S(r['Diagnosis']), act if insp != 'OK' else '', '', stage, 'WearCheck laboratory', '', '', rem))
VIBST = {'NORMAL':'OK', 'BORDERLINE':'Alert', 'URGENT':'Danger', 'CRITICAL':'Danger'}
for _, r in vib.iterrows():
    insp = VIBST.get(S(r['Final Severity']).upper(), ''); dte = D(r['Date Measured']); dl = D(r['Action Deadline'])
    ast = S(r['Action Status']); wo = S(r['WO / WR Number']); closed = verified = ''
    if insp == 'OK' or not insp: stage = 'No action'
    elif ast.lower() == 'completed':
        stage = 'Closed'; closed = (dl if dl and dl >= dte else dte).strftime('%Y-%m-%d'); verified = 'Not recorded'; n_est += 1
    else: stage = 'WO raised' if wo else 'Raised'
    num = lambda c: float(r[c]) if not pd.isna(r[c]) else ''
    rem = f"Failure mode: {S(r['Failure Mode'])}; bearing {S(r['Bearing Condition gE'])} gE; {S(r['Data Source'])}" + ('; close date estimated from action deadline' if closed else '')
    recs.append((dte.strftime('%Y-%m-%d'), r['tag'], 'Vibration & visual route', 'Vibration', f"{S(r['Measurement Point'])} {S(r['Direction'])} overall velocity".strip(),
                 num('Overall Velocity mm/s RMS'), 'mm/s', num('Alarm Limit mm/s'), num('Trip Limit mm/s'), 'Higher is worse', insp,
                 S(r['Diagnosis']), S(r['Recommended Action']) if insp != 'OK' else '', wo, stage, S(r['Analyst']), closed, verified, rem))
recs.extend(VREC)
recs.sort(key=lambda x: (x[0], x[1]))
# ---- Professional coding of every record: ISO 14224 failure mode / mechanism, API 571 damage mechanism,
#      execution window and likelihood override (risk matrix). Rules are documented in docs/kpi-definitions.md.
def _enrich(x):
    x = tuple(x) + ('',) * (19 - len(x)) if len(x) < 19 else tuple(x)
    src, tech, par, insp, fnd, rem = x[2], x[3], x[4], x[10], (x[11] or ''), (x[18] or '')
    bad = insp in ('Alert', 'Danger')
    fm = mech = dm = win = lik = ''
    m = re.search(r'Failure mode: ([^;]+)', rem)
    if src == 'Vibration & visual route' and bad:
        fm, mech = AS.site_fm(m.group(1) if m else '')
        if tech == 'Vibration' and not fm: fm, mech = 'VIB', '1.2'
    elif src == 'Oil analysis' and bad:
        fm, mech = 'PDE', '5.2'
    elif src == 'Statutory inspection' and bad and par != 'External visual inspection (overall)':
        t = (rem + ' ' + fnd).lower(); leak = ('leak' in t) or ('loss of containment' in t)
        fm, mech = ('ELP', '1.1') if leak else ('STD', '2.2')
        dm = AS.mech_from_text(rem) or AS.mech_from_text(fnd)
        win = 'Unit isolation' if any(k in t for k in ['gasket', 'bolt', 'flange', 'manway', 'roof', 'foundation', 'concrete', 'anchor', 'nozzle', 'shell']) else 'Online'
        lik = 5 if 'active loss of containment' in t else ''
    elif src == 'Ultrasonic report':
        if par.startswith('Shell band'): dm = 'Wall thinning (general)'
        else:
            cat = par.split(':')[0]; t = fnd.lower()
            fm, mech, dm, win, lik = {
                'Thickness': ('STD', '2.2', 'Wall thinning (general)', 'Online', ''),
                'Through-wall': ('ELP', '2.2', 'Wall thinning (general)', 'Unit isolation', 5),
                'Weld': ('ELP' if 'leak' in t else 'STD', '2.0', 'Weld cracking / weld defects', 'Unit isolation', 5 if 'leak' in t else ''),
                'Coating': ('STD', '2.2', 'Coating breakdown', 'Online', ''),
                'Lining': ('STD', '2.2', 'Lining failure', 'Unit isolation', ''),
            }.get(cat, ('STD', '1.4', 'Foundation settlement', 'Unit isolation', '') if ('foundation' in t or 'lean' in t) else ('STD', '2.2', 'Atmospheric corrosion', 'Online', ''))
    return x[:19] + (fm, mech, dm, win, lik)
recs = [_enrich(x) for x in recs]


# ---- Events and schedule from the Pronto WO export
wo = pd.read_excel(SRC, 'WO Data', dtype={'Work Order': str}); wo = wo[wo['Work Order'].notna()].copy()
NAME2TAG = {}
for a in assets: NAME2TAG.setdefault(a[2].upper().strip(), []).append(a[0])
tagset = set(tags)
def match(pi):
    p = S(pi); m = re.match(r'^(\d{3}-[A-Z0-9]+-[A-Z0-9]+(?:[ -][A-Z0-9]+)?)', p.upper())
    if m:
        for cand in [m.group(1), m.group(1).split(' ')[0]]:
            if cand in tagset: return cand
    c = NAME2TAG.get(p.upper())
    return c[0] if c and len(c) == 1 else None
ev = []; um = 0
for _, r in wo[wo['Work Type'].isin(['Breakdown', 'Corrective Maintenance']) & (wo['Status'] != 'Cancelled')].iterrows():
    t = match(r['Plant Item'])
    if not t: um += 1; continue
    dte = D(r['Actual Finish Date']) or D(r['Est. Start Date'])
    dn = r['Actual Down Time']; dn = float(dn) if not pd.isna(dn) else 0
    ev.append((dte.strftime('%Y-%m-%d'), t, 'Failure' if r['Work Type'] == 'Breakdown' else 'Repair', S(r['Work Description']), dn if r['Work Type'] == 'Breakdown' else 0, S(r['Work Order']), ''))
# ---- Pronto WO export (WO Data) -> Work_Orders input (all work types; merged by WO no. on every import)
def _wtype(t, scope):
    t = (t or '').lower()
    if 'prevent' in t or 'statut' in t: return 'Preventive'
    if 'breakdown' in t: return 'Breakdown'
    if 'emerg' in t: return 'Emergency'
    if 'shutdown' in t: return 'Shutdown'
    if any(k in t for k in ('project', 'improv', 'modif')): return 'Improvement'
    if 'condition' in t or 'inspect' in t: return 'Condition-based'
    return 'Corrective'
_WST = {'complete': 'Complete', 'completed': 'Complete', 'in progress': 'In progress', 'planned': 'Scheduled', 'scheduled': 'Scheduled',
        'released': 'Ready', 'forecast': 'Open', 'on hold': 'Open', 'open': 'Open', 'cancelled': 'Cancelled'}
def _num(v):
    v = pd.to_numeric(v, errors='coerce'); return None if pd.isna(v) else float(v)
PRONTO_WO = []
for _, r in wo.iterrows():
    st = _WST.get(S(r.get('Status')).lower(), 'Open')
    sched = D(r.get('Scheduled')) or D(r.get('Est. Start Date'))
    PRONTO_WO.append([S(r['Work Order']), match(r['Plant Item']) or '', S(r.get('Work Description'))[:120], _wtype(S(r.get('Work Type')), S(r.get('Scope'))),
                      S(r.get('Priority')), st, S(r.get('Responsibility')) or S(r.get('Section')), D(r.get('Est. Start Date')),
                      D(r.get('Required')) or D(r.get('Latest')) or D(r.get('Est. Finish Date')), sched, D(r.get('Actual Finish Date')) or (D(r.get('Finish Date')) if st == 'Complete' else None),
                      _num(r.get('Estimated Hours')), _num(r.get('Actual Hours')), _num(r.get('Actual Down Time')) or 0, S(r.get('Fault Code 1')) or S(r.get('Fault 1')),
                      _num(r.get('Actual Cost')), '', '', ''])
rv('Summary', 'Pronto work orders', f'{len(PRONTO_WO)} work orders mapped to Work_Orders ({sum(1 for x in PRONTO_WO if x[1])} linked to an asset). Raised date = Pronto Est. Start Date (no creation date in the export).')

TECHMAP = {'Lubrication (Greasing)':'Lubrication / oil', 'Oil Change/Top-up':'Lubrication / oil', 'Oil Sampling/Analysis':'Lubrication / oil',
           'Tank/Vessel Inspection':'Tank & vessel', 'Statutory Inspection (Other)':'Statutory'}
cm = wo[wo['Scope'].notna() & (wo['Status'] != 'Cancelled')].copy()
cm['tech'] = cm['Task Category'].map(lambda x: TECHMAP.get(S(x), 'Visual'))
cm['month'] = pd.to_datetime(cm['Est. Start Date']).dt.to_period('M').dt.to_timestamp()
sch = [(m.strftime('%Y-%m-%d'), t, int(len(g)), int((g['Status'] == 'Complete').sum())) for (m, t), g in cm.groupby(['month', 'tech'])]
wo_months = sorted({x[0][:7] for x in sch})
# ---- Pronto work-order summary per month and work type (SMRP metrics); months already stored are kept
_wo = wo.copy()
_wo['month'] = pd.to_datetime(_wo['Est. Start Date'], errors='coerce').dt.to_period('M').dt.to_timestamp()
_wo = _wo[_wo['month'].notna() & (_wo['Status'] != 'Cancelled')]
WOSUM = []
for (m_, wt), g in _wo.groupby(['month', 'Work Type']):
    WOSUM.append([m_.to_pydatetime(), S(wt), int(len(g)), int((g['Status'] == 'Complete').sum()), int((g['Status'] == 'In Progress').sum()),
                  int(g['Status'].isin(['Forecast', 'Planned', 'On Hold']).sum())])
_new_months = {r[0].strftime('%Y-%m') for r in WOSUM}
for r in (AS.old_sheet_rows(OUT, 'WO_Summary', 6) or []):
    if r[0] is not None and pd.Timestamp(r[0]).strftime('%Y-%m') not in _new_months: WOSUM.append(r)
WOSUM.sort(key=lambda r: (pd.Timestamp(r[0]), str(r[1])))
rv('Summary', 'Pronto WO summary', f"{len(WOSUM)} month x work-type rows (months: {', '.join(sorted({pd.Timestamp(r[0]).strftime('%b %Y') for r in WOSUM}))}) for SMRP metrics")


rv('Summary', 'Source', os.path.basename(SRC))
rv('Summary', 'Data entry rows', f'{len(de)} inspections, {de["date"].min():%d-%b-%Y} to {MAXD:%d-%b-%Y}')
rv('Summary', 'Assets in register', f'{len(assets)} ({len(MASTER)} tags in Criticality sheet)')
rv('Summary', 'Records created', f'{len(recs)} ({sum(1 for x in recs if x[3]=="Visual")} visual, {sum(1 for x in recs if x[3]=="Lubrication / oil")} oil, {sum(1 for x in recs if x[3]=="Vibration")} vibration, {sum(1 for x in recs if x[2]=="Statutory inspection")} vessel statutory, {sum(1 for x in recs if x[2]=="Ultrasonic report")} UT)')
rv('Summary', 'Offline inspections excluded', f'{n_off} rows: equipment offline, no condition assessed')
rv('Summary', 'Rows without status excluded', n_nostat + bad_dates)
rv('Summary', 'Alert/Danger kept as readings only', f'{n_untracked} rows with record status blank or Observation (listed below). To track one, set its Record Status to Open in the CM workbook Data entry and re-import (Records are regenerated on each import).')
rv('Summary', 'Closed dates', f'{n_rec} recorded in the CM workbook (Date Closed column); {n_est} estimated (next inspection date or action deadline) because Date Closed is blank')
rv('Summary', 'Criticality conflicts', f'{len(conflicts)} assets where Data entry differs from the Criticality sheet (listed below)')
rv('Summary', 'Pronto WO export', f'Months: {", ".join(wo_months)}. {len(ev)} breakdown/corrective WOs linked to assets, {um} could not be linked (plant item is a name, not a tag)')
rv('Method', 'ACI factors', 'Safety, production, redundancy and repair cost set from site criticality (C1 = 5,5,4,4 · C2 = 3,4,3,2 · C3 = 2,2,2,1 · C4 = 1,1,1,1). Failure history from Alert/Danger findings in the last 12 months (0 = 1 … 6+ = 5). Refine in a criticality workshop.')
rv('Method', 'Area names', 'Area groups and names are provisional, inferred from equipment names. Correct them in Asset_Register columns D and E (or in AREA_NAMES in scripts/import_cm_workbook.py).')
rv('Method', 'Pronto asset no.', 'Equipment ID is used as the asset key. Replace with the Pronto asset number if different.')
for k in ['Statutory certificates', 'CM value / cost avoidance', 'Programme cost', 'Management decisions', 'Fleet prestarts']:
    rv('Not in source', k, 'No data in the CM workbook. Sheet left empty: fill it to activate the related dashboard panel.')
rv('Not in source', 'Reliability history', 'Pronto WO export covers ' + ', '.join(wo_months) + ' only, so MTBF/MTTR and bad actors need a longer WO export.')
# ---- Draft commentary for the latest month (only if none exists)
def _draft():
    M = MAXD.to_period('M'); P = M - 1
    cur = de[de['date'].dt.to_period('M') == M]; prv = de[de['date'].dt.to_period('M') == P]
    off = (cur['Status'] == 'Offline').sum(); online = len(cur) - off
    na = (cur['Status'] == 'Alert').sum(); nd = (cur['Status'] == 'Danger').sum()
    fm = cur[cur['Status'].isin(['Alert', 'Danger'])]['Failure Mode'].value_counts()
    fm = ', '.join(f'{k.lower()} ({v})' for k, v in list((k, v) for k, v in fm.items() if k not in ('Other', 'No defect', 'Equipment offline'))[:5]) or 'none recorded'
    opn = [x for x in recs if x[14] in ('Raised', 'WO raised', 'Scheduled', 'Awaiting verification')]
    old = sum(1 for x in opn if (MAXD - pd.Timestamp(x[0])).days > 30)
    nwo = sum(1 for x in opn if x[13])
    ch = (f"{len(cur)} inspections in {M.strftime('%b %Y')} ({online} online, {off} offline = {round(off / max(len(cur), 1) * 100)}%) vs {len(prv)} in {P.strftime('%b %Y')}. "
          f"{na} Alert and {nd} Danger findings raised this month.")
    why = f"Main failure modes this month: {fm}. [Analyst: add the drivers, e.g. shutdown, access, route resourcing, recurring defects.]"
    do = (f"{len(opn)} recommendations open, {old} older than 30 days, {nwo} with a work order. "
          f"Weekly CM / planner review to raise WOs and close verified findings. [Analyst: add specific actions and dates.]")
    return [M.to_timestamp().to_pydatetime(), ch, why, do, 'DRAFT: auto-generated, review before presenting']
_cm = OLD.get('Commentary', [])
_lm = MAXD.to_period('M')
if not any(pd.Timestamp(r[0]).to_period('M') == _lm for r in _cm if r[0] is not None):
    _cm = _cm + [_draft()]
    rv('Summary', 'Commentary', f"Draft commentary written for {_lm.strftime('%b %Y')}: review and edit the Commentary sheet before presenting")

print(f'Extracted {len(assets)} assets, {len(recs)} records, {len(ev)} events, {len(sch)} schedule rows')

F='Arial'
NAVY='203646'; GREY='5D6E79'; INFILL='FFF8DC'; FFILL='EEF2F4'
thin=Side(style='thin',color='C9D3D9'); B=Border(left=thin,right=thin,top=thin,bottom=thin)
def font(**k): k.setdefault('name',F); k.setdefault('size',10); return Font(**k)
def fill(c): return PatternFill('solid',start_color=c,end_color=c)
OKF,ALF,DGF=fill('D7EEDF'),fill('FBEBC0'),fill('F4CFC9')

wb=Workbook()
G=wb.active; G.title='Guide'
AR=wb.create_sheet('Asset_Register'); RC=wb.create_sheet('Records'); LS=wb.create_sheet('Lists')

# ---------------- LISTS ----------------
lists={
 'A':('Areas',AREA_LIST),
 'B':('Asset classes',TYPE_LIST),
 'C':('Sources (input)',['Vibration & visual route','IR survey','Substation inspection','Statutory inspection','Ultrasonic report','Oil analysis','DG checklist','LV/HV inspection checklist','LV/HV prestart checklist','Ad hoc / operator report']),
 'D':('Techniques',['Vibration','Lubrication / oil','Ultrasonics (airborne)','Ultrasonic thickness','Infrared','Visual','Statutory','Tank & vessel','Piping','Structural','Electrical','DG checks','Fleet inspection','Prestart']),
 'E':('Record stages',['No action','Raised','WO raised','Scheduled','Awaiting verification','Closed']),
 'F':('Status',['OK','Alert','Danger']),
 'G':('Limit direction',['Higher is worse','Lower is worse']),
 'H':('Asset state',['Active','Standby','Out of service','Decommissioned']),
 'I':('Operating mode',['Duty','Standby','Mobile']),
 'J':('Yes / No',['Y','N']),
}
for c,(h,items) in lists.items():
    LS[f'{c}1']=h; LS[f'{c}1'].font=font(bold=True,color='FFFFFF'); LS[f'{c}1'].fill=fill(NAVY)
    for i,v in enumerate(items): LS[f'{c}{i+2}']=v; LS[f'{c}{i+2}'].font=font()
    LS.column_dimensions[c].width=max(16,max(len(x) for x in items+[h])+2)
def ptab(r,hdr,rows,note=None):
    for j,h in enumerate(hdr):
        c=LS.cell(r,12+j,h); c.font=font(bold=True,color='FFFFFF'); c.fill=fill(NAVY)
    for i,row in enumerate(rows):
        for j,v in enumerate(row):
            c=LS.cell(r+1+i,12+j,v); c.font=font(color='0000FF' if j>0 else '000000'); c.border=B
            if j>0: c.fill=fill(INFILL)
    if note: LS.cell(r+1+len(rows),12,note).font=font(italic=True,color=GREY,size=9)
ptab(1,['ACI factor','Weight'],[['Safety & environment',0.30],['Production impact',0.30],['Redundancy',0.15],['Repair cost / MTTR',0.15],['Failure history',0.10]],'Weights must total 1.00. Agree with Engineering Manager.')
ptab(9,['Priority','Minimum score','Response time (days)'],[['P1',70,7],['P2',45,30],['P3',0,60]])
ptab(14,['Severity','Factor'],[['Danger',1.0],['Alert',0.6]],'Priority score = ACI x severity factor. Any Danger is always P1.')
ptab(19,['Criticality class','Minimum ACI'],[['A',70],['B',45],['C',0]])
ptab(23,['Statutory rule','Days'],[['Expiry alert window',60]],'Certificate expiring within this window = Alert; expired = Danger')
ptab(27,['Condition score (for AHI)','Score'],[['OK',100],['Alert',60],['Danger',20],['Unknown',50]],'Unknown = not inspected within the required interval')
LS.column_dimensions['L'].width=28; LS.column_dimensions['M'].width=15; LS.column_dimensions['N'].width=20
LS.cell(52,12,'Blue text on cream = settings you can change. Every formula in the workbook reads from these cells.').font=font(italic=True,color=GREY,size=9)

def rng(c): n=len(lists[c][1]); return f'=Lists!${c}$2:${c}${n+1}'

# ---------------- ASSET REGISTER ----------------
TECH=lists['D'][1]
reg_cols=[('Pronto asset no.',14,'in'),('Site tag',13,'in'),('Description',28,'in'),('Area',13,'in'),('System / location',20,'in'),('Asset class',22,'in'),
 ('Manufacturer',14,'in'),('Model',12,'in'),('Serial no.',12,'in'),('Operating mode',11,'in'),('Asset state',11,'in'),
 ('Safety & env. (1-5)',9,'in'),('Production (1-5)',9,'in'),('Redundancy (1-5)',9,'in'),('Repair cost / MTTR (1-5)',9,'in'),('Failure history (1-5)',9,'in'),
 ('ACI (0-100)',8,'f'),('Class',7,'f'),('Current status',10,'f'),('Open records',8,'f')]
reg_cols+= [(t,6.5,'in') for t in TECH]
reg_cols+= [('Statutory cert. required',9,'in'),('Certificate expiry',12,'in'),('Days to expiry',9,'f'),('Statutory status',10,'f'),('Responsible',14,'in'),('Remarks',30,'in')]
reg_cols+= [(h,w,'in') for h,w in AS.REG_EXTRA]  # cols 41-45
RN=1000; RR=3000  # register rows, record rows
def header(ws,cols,title,sub):
    ws['A1']=title; ws['A1'].font=font(bold=True,size=14,color=NAVY)
    ws['A2']=sub; ws['A2'].font=font(size=9,color=GREY)
    for j,(h,w,k) in enumerate(cols,1):
        c=ws.cell(4,j,h); c.font=font(bold=True,color='FFFFFF' if k=='in' else '1D2A32',size=9)
        c.fill=fill(NAVY if k=='in' else 'C9D3D9'); c.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center'); c.border=B
        ws.column_dimensions[L(j)].width=w
    ws.row_dimensions[4].height=48
header(AR,reg_cols,'AHIM Master Asset Register','Dark headers = enter data. Grey headers = automatic (do not type). Imported from the CM Monthly Report: see Import_Review. Pronto asset no. is the unique key used everywhere.')
tc0=21  # first technique column (U)
# technique group label
AR.cell(3,tc0,'Applicable techniques (enter Y)').font=font(bold=True,size=9,color=NAVY)
for j in range(tc0,tc0+len(TECH)):
    AR.cell(4,j).alignment=Alignment(text_rotation=90,horizontal='center',vertical='bottom')
AR.row_dimensions[4].height=110
AR.cell(3,12,'Criticality factors: 1 = low impact, 5 = severe').font=font(bold=True,size=9,color=NAVY)

T=dict(VA='Vibration',OA='Lubrication / oil',US='Ultrasonics (airborne)',UT='Ultrasonic thickness',IR='Infrared',VI='Visual',ST_='Statutory',TV='Tank & vessel',PI='Piping',STR='Structural',EI='Electrical',DG='DG checks',FI='Fleet inspection',PS='Prestart')
# assets built from the CM workbook above
tkeys=list(T.keys())
SAMPLE=font()
for i,a in enumerate(assets):
    r=5+i
    vals=[a[0],a[1],a[2],a[3],a[4],a[5],a[6],a[7],'',a[8],'Active']+a[9]
    for j,v in enumerate(vals,1): AR.cell(r,j,v)
    for k in a[10].split(): AR.cell(r,tc0+tkeys.index(k),'Y')
    AR.cell(r,35,a[11])
    if a[12]: AR.cell(r,36,dt.datetime.strptime(a[12],'%Y-%m-%d'))
    AR.cell(r,39,a[13])
    _x=AS.merge_extras(a, OLDX.get(a[0],{}) if 'OLDX' in globals() else {})
    for _k,_v in enumerate(_x):
        if _v not in (None,''): AR.cell(r,41+_k,_v)
    if not _x[0] and 'rv' in globals(): rv('Strategy class to confirm', a[0], f"{a[2]}: equipment type not identified from the source data; set Strategy class in Asset_Register")
NA=len(assets)
for r in range(5,RN+1):
    f=lambda col: f'${col}{r}'
    AR.cell(r,17,f'=IF(COUNT($L{r}:$P{r})<5,"",ROUND(($L{r}*Lists!$M$2+$M{r}*Lists!$M$3+$N{r}*Lists!$M$4+$O{r}*Lists!$M$5+$P{r}*Lists!$M$6)/5*100,0))')
    AR.cell(r,18,f'=IF($Q{r}="","",IF($Q{r}>=Lists!$M$20,"A",IF($Q{r}>=Lists!$M$21,"B","C")))')
    AR.cell(r,19,f'=IF($A{r}="","",IF(COUNTIF(Records!$C$5:$C${RR},$A{r})=0,"No data",CHOOSE(_xlfn.MAXIFS(Records!$R$5:$R${RR},Records!$C$5:$C${RR},$A{r},Records!$Y$5:$Y${RR},"<>Closed")+1,"OK","Alert","Danger")))')
    AR.cell(r,20,f'=IF($A{r}="","",COUNTIFS(Records!$C$5:$C${RR},$A{r},Records!$R$5:$R${RR},">0",Records!$Y$5:$Y${RR},"<>Closed"))')
    AR.cell(r,37,f'=IF($AJ{r}="","",$AJ{r}-TODAY())')
    AR.cell(r,38,f'=IF($AJ{r}="","",IF($AK{r}<0,"Danger",IF($AK{r}<=Lists!$M$24,"Alert","OK")))')
    for j in range(1,len(reg_cols)+1):
        c=AR.cell(r,j); c.border=B
        kind=reg_cols[j-1][2]
        if kind=='f': c.fill=fill(FFILL); c.font=font()
        else: c.font=SAMPLE if r<5+NA else font()
        if j>=12 and j!=36 and j<=38 or (tc0<=j<tc0+len(TECH)): c.alignment=Alignment(horizontal='center')
    AR.cell(r,36).number_format='dd-mmm-yy'
AR.freeze_panes='D5'
AR.auto_filter.ref=f'A4:{L(len(reg_cols))}{RN}'
def dv(ws,formula,ref,prompt=None,typ='list',**k):
    d=DataValidation(type=typ,formula1=formula,allow_blank=True,**k)
    if prompt: d.promptTitle='Input'; d.prompt=prompt; d.showInputMessage=True
    d.error='Choose a value from the list.'; d.showErrorMessage=True
    ws.add_data_validation(d); d.add(ref)
dv(AR,rng('A'),f'D5:D{RN}'); dv(AR,rng('B'),f'F5:F{RN}'); dv(AR,rng('I'),f'J5:J{RN}'); dv(AR,rng('H'),f'K5:K{RN}')
dv(AR,'1',f'L5:P{RN}',typ='whole',operator='between',formula2='5',prompt='Score 1 (low impact) to 5 (severe)')
dv(AR,'"Y"',f'{L(tc0)}5:{L(tc0+len(TECH)-1)}{RN}')
dv(AR,rng('J'),f'AI5:AI{RN}')
dv(AR,'DATE(2000,1,1)',f'AJ5:AJ{RN}',typ='date',operator='greaterThan')
# duplicate key check
AR.conditional_formatting.add(f'A5:A{RN}',FormulaRule(formula=[f'AND($A5<>"",COUNTIF($A$5:$A${RN},$A5)>1)'],fill=DGF))
for col in ['S','AL']:
    AR.conditional_formatting.add(f'{col}5:{col}{RN}',CellIsRule(operator='equal',formula=['"Danger"'],fill=DGF,font=font(bold=True,color='9C1C10')))
    AR.conditional_formatting.add(f'{col}5:{col}{RN}',CellIsRule(operator='equal',formula=['"Alert"'],fill=ALF,font=font(bold=True,color='7A5300')))
    AR.conditional_formatting.add(f'{col}5:{col}{RN}',CellIsRule(operator='equal',formula=['"OK"'],fill=OKF,font=font(color='1F6B40')))
AR.conditional_formatting.add(f'R5:R{RN}',CellIsRule(operator='equal',formula=['"A"'],font=font(bold=True,color=NAVY)))
AR['A4'].comment=Comment('Unique key. Use the asset number exactly as in Pronto so records link to work orders. Duplicates turn red.','AHIM')
AR['S4'].comment=Comment('Worst status among this asset\'s records that are not Closed. "No data" = no record yet.','AHIM')

# ---------------- RECORDS ----------------
rec_cols=[('Record ID',10,'f'),('Date',11,'in'),('Pronto asset no.',13,'in'),('Description',24,'f'),('Area',12,'f'),('Class',6,'f'),('ACI',6,'f'),
 ('Source (input)',22,'in'),('Technique',18,'in'),('Component / parameter / checklist item',30,'in'),('Value',8,'in'),('Unit',8,'in'),
 ('Alert limit',8,'in'),('Danger limit',8,'in'),('Limit direction',13,'in'),('Inspector status',10,'in'),('Status',9,'f'),('Severity',7,'f'),
 ('Finding',36,'in'),('Recommendation',36,'in'),('Priority score',8,'f'),('Priority',8,'f'),('Due date',11,'f'),('Pronto WO no.',11,'in'),
 ('Record stage',15,'in'),('Inspector',14,'in'),('Date closed',11,'in'),('Verified by',14,'in'),('Days open',7,'f'),('Due status',13,'f'),('Remarks',28,'in')]
rec_cols+= [(h,w,'in') for h,w in AS.REC_EXTRA]  # cols 32-36
header(RC,rec_cols,'AHIM Common Record','One row per reading, checklist defect or certificate check, from ANY source. Dark headers = enter. Grey = automatic. Imported from the CM Monthly Report. New rows can be typed below.')
d=lambda s: dt.datetime.strptime(s,'%Y-%m-%d')
# records built from the CM workbook above


NR=len(recs)
for i,x in enumerate(recs):
    r=5+i
    RC.cell(r,2,d(x[0])); RC.cell(r,3,x[1]); RC.cell(r,8,x[2]); RC.cell(r,9,x[3]); RC.cell(r,10,x[4])
    for col,v in zip([11,12,13,14,15,16],x[5:11]):
        if v!='': RC.cell(r,col,v)
    RC.cell(r,19,x[11]); RC.cell(r,20,x[12] or None); RC.cell(r,24,x[13] or None); RC.cell(r,25,x[14]); RC.cell(r,26,x[15])
    if x[16]: RC.cell(r,27,d(x[16]))
    if x[17]: RC.cell(r,28,x[17])
    if len(x)>18 and x[18]: RC.cell(r,31,x[18])
    for _k,_v in enumerate(x[19:24]):
        if _v not in (None,''): RC.cell(r,32+_k,_v)
REG=f'Asset_Register!$A$5:$A${RN}'
def lk(col,r): return f'=IF($C{r}="","",IFERROR(INDEX(Asset_Register!${col}$5:${col}${RN},MATCH($C{r},{REG},0)),"Not in register"))'
for r in range(5,RR+1):
    RC.cell(r,1,f'=IF($C{r}="","","R"&TEXT(ROW()-4,"00000"))')
    RC.cell(r,4,lk('C',r)); RC.cell(r,5,lk('D',r)); RC.cell(r,6,lk('R',r)); RC.cell(r,7,lk('Q',r))
    RC.cell(r,17,f'=IF($C{r}="","",IF($P{r}<>"",$P{r},IF(AND(ISNUMBER($K{r}),ISNUMBER($M{r}),ISNUMBER($N{r})),IF($O{r}="Lower is worse",IF($K{r}<=$N{r},"Danger",IF($K{r}<=$M{r},"Alert","OK")),IF($K{r}>=$N{r},"Danger",IF($K{r}>=$M{r},"Alert","OK"))),"Not set")))')
    RC.cell(r,18,f'=IF(OR($Q{r}="",$Q{r}="Not set"),"",MATCH($Q{r},Lists!$F$2:$F$4,0)-1)')
    RC.cell(r,21,f'=IF(OR($Q{r}="",$Q{r}="OK",$Q{r}="Not set",NOT(ISNUMBER($G{r}))),"",ROUND($G{r}*IF($Q{r}="Danger",Lists!$M$15,Lists!$M$16),0))')
    RC.cell(r,22,f'=IF($U{r}="","",IF(OR($Q{r}="Danger",$U{r}>=Lists!$M$10),"P1",IF($U{r}>=Lists!$M$11,"P2","P3")))')
    RC.cell(r,23,f'=IF(OR($V{r}="",$B{r}=""),"",$B{r}+INDEX(Lists!$N$10:$N$12,MATCH($V{r},Lists!$L$10:$L$12,0)))')
    RC.cell(r,29,f'=IF(OR($U{r}="",$B{r}=""),"",IF($AA{r}<>"",$AA{r}-$B{r},TODAY()-$B{r}))')
    RC.cell(r,30,f'=IF($W{r}="","",IF($Y{r}="Closed",IF($AA{r}="","Closed",IF($AA{r}<=$W{r},"Closed on time","Closed late")),IF(TODAY()>$W{r},"Overdue","Within time")))')
    for j in range(1,len(rec_cols)+1):
        c=RC.cell(r,j); c.border=B
        if rec_cols[j-1][2]=='f': c.fill=fill(FFILL); c.font=font()
        else: c.font=SAMPLE if r<5+NR else font()
        if j in (19,20,31,10): c.alignment=Alignment(wrap_text=True,vertical='top')
    for j in (2,23,27): RC.cell(r,j).number_format='dd-mmm-yy'
RC.freeze_panes='D5'
RC.auto_filter.ref=f'A4:{L(len(rec_cols))}{RR}'
dv(RC,'DATE(2000,1,1)',f'B5:B{RR}',typ='date',operator='greaterThan',prompt='Date of inspection or reading')
dv(RC,f'={REG}',f'C5:C{RR}',prompt='Pick the Pronto asset no. from the register')
dv(RC,rng('C'),f'H5:H{RR}'); dv(RC,rng('D'),f'I5:I{RR}'); dv(RC,rng('G'),f'O5:O{RR}')
dv(RC,rng('F'),f'P5:P{RR}',prompt='Required for checklists, visual and statutory. For measured values leave blank (status is calculated) or set to override.')
dv(RC,rng('E'),f'Y5:Y{RR}')
for col in ['Q','P']:
    RC.conditional_formatting.add(f'{col}5:{col}{RR}',CellIsRule(operator='equal',formula=['"Danger"'],fill=DGF,font=font(bold=True,color='9C1C10')))
    RC.conditional_formatting.add(f'{col}5:{col}{RR}',CellIsRule(operator='equal',formula=['"Alert"'],fill=ALF,font=font(bold=True,color='7A5300')))
    RC.conditional_formatting.add(f'{col}5:{col}{RR}',CellIsRule(operator='equal',formula=['"OK"'],fill=OKF,font=font(color='1F6B40')))
RC.conditional_formatting.add(f'Q5:Q{RR}',CellIsRule(operator='equal',formula=['"Not set"'],font=font(italic=True,color='9C1C10')))
RC.conditional_formatting.add(f'V5:V{RR}',CellIsRule(operator='equal',formula=['"P1"'],fill=fill('BE3528'),font=font(bold=True,color='FFFFFF')))
RC.conditional_formatting.add(f'AD5:AD{RR}',CellIsRule(operator='equal',formula=['"Overdue"'],font=font(bold=True,color='BE3528')))
RC.conditional_formatting.add(f'AD5:AD{RR}',CellIsRule(operator='equal',formula=['"Closed late"'],font=font(color='B57C00')))
RC.conditional_formatting.add(f'D5:D{RR}',CellIsRule(operator='equal',formula=['"Not in register"'],fill=DGF))
RC.conditional_formatting.add(f'Y5:Y{RR}',FormulaRule(formula=[f'AND(OR($Q5="Alert",$Q5="Danger"),$Y5="")'],fill=ALF))
RC['P4'].comment=Comment('Checklists / visual / statutory: choose OK, Alert or Danger using the severity rules in the Guide. Measured values: leave blank and status is calculated from the limits; fill only to override (analyst judgement).','AHIM')
RC['Y4'].comment=Comment('OK records: No action. Alert/Danger: Raised -> WO raised -> Scheduled -> Awaiting verification -> Closed. Close only after a re-check confirms the fix.','AHIM')

# ---------------- GUIDE ----------------
G.sheet_view.showGridLines=False
G.column_dimensions['A'].width=2
for c,w in zip('BCDEFG',[26,26,24,26,18,30]): G.column_dimensions[c].width=w
row=[1]
def put(txt,**k):
    c=G.cell(row[0],2,txt); c.font=font(**k); c.alignment=Alignment(wrap_text=False); row[0]+=1; return c
def para(txt):
    G.merge_cells(start_row=row[0],start_column=2,end_row=row[0],end_column=7)
    c=G.cell(row[0],2,txt); c.font=font(); c.alignment=Alignment(wrap_text=True,vertical='top')
    G.row_dimensions[row[0]].height=15*max(1,len(txt)//120+1); row[0]+=1
def table(hdr,rows):
    for j,h in enumerate(hdr):
        c=G.cell(row[0],2+j,h); c.font=font(bold=True,color='FFFFFF'); c.fill=fill(NAVY); c.alignment=Alignment(wrap_text=True,vertical='center'); c.border=B
    row[0]+=1
    for rw in rows:
        mx=1
        for j,v in enumerate(rw):
            c=G.cell(row[0],2+j,v); c.font=font(); c.alignment=Alignment(wrap_text=True,vertical='top'); c.border=B
            w=G.column_dimensions['BCDEFG'[j]].width; mx=max(mx,len(str(v))//int(w*1.1)+1)
        G.row_dimensions[row[0]].height=13.5*mx+2; row[0]+=1
    row[0]+=1
put('AHIM: Master Asset Register & Common Record',bold=True,size=16,color=NAVY)
put('Asset Health & Integrity Management · Lotus Africa Uranium Plant',size=10,color=GREY); row[0]+=1
put('Purpose',bold=True,size=12,color=NAVY)
para('Every inspection input (vibration and visual routes, IR, substation, statutory, ultrasonic, oil analysis, diesel generator checklists, light vehicle and heavy equipment inspections and prestarts) is entered in ONE format against ONE asset list. The register and records then feed the AHIM dashboard: plant health, criticality, top priorities, open records and machine history.')
row[0]+=1
put('How it works',bold=True,size=12,color=NAVY)
table(['Step','What to do','Who'],[
 ['1. Register the asset','Add it to Asset_Register with its Pronto asset no., area, class, criticality factors and applicable techniques. ACI and class calculate automatically.','RCM Specialist (once per asset)'],
 ['2. Enter the result','Each reading, checklist defect or certificate check becomes one row in Records. Pick the asset from the dropdown; description, area and ACI fill in.','Input owner (analyst, electrician, DG operator, fleet mechanic, contractor)'],
 ['3. Status is set','Measured values: from the Alert/Danger limits. Checklists, visual, statutory: Inspector status using the severity rules below.','Automatic / inspector'],
 ['4. Priority and due date','Alert and Danger rows get a priority score (ACI x severity), P1/P2/P3 and a due date automatically. Any Danger is always P1, whatever the asset criticality.','Automatic'],
 ['5. Action and close','Raise the Pronto WO, enter the WO no., move the stage forward. Close only after a re-check confirms the fix, and record who verified it.','RCM Specialist / planner'],
 ['6. Review','The asset\'s Current status in the register always shows the worst open record. Use filters on Records for overdue items.','RCM Specialist (weekly)'],
])
put('Colour legend',bold=True,size=12,color=NAVY)
table(['Item','Meaning'],[['Dark blue header','Enter data in this column'],['Grey header, grey cells','Automatic formula: do not type over'],['Import_Review sheet','What was imported, assumptions and items to check'],['Blue on cream (Lists sheet)','Settings: weights, thresholds, response times'],['Red key cell (register)','Duplicate Pronto asset no.'],['Amber Record stage','Alert/Danger row with no stage yet']])
put('How each input source maps into the record',bold=True,size=12,color=NAVY)
table(['Source (input)','Technique(s)','Record type','How status is set','Typical frequency','Owner'],[
 ['Vibration & visual route','Vibration, Visual','Measured + checklist','Limits (ISO 20816-3: 4.5 / 7.1 mm/s); visual by severity rules','Monthly route','RCM Specialist'],
 ['IR survey','Infrared','Measured','dT limits (NETA: >4 °C Alert, >15 °C Danger) or absolute temperature limits','Monthly','RCM Specialist'],
 ['Substation inspection','Infrared, Visual, Electrical','Measured + checklist','Limits for readings; severity rules for checklist items','Monthly','Electrical'],
 ['Statutory inspection','Statutory','Compliance','Inspector status; register also flags expiry (60 days = Alert, expired = Danger)','Per certificate','RCM Specialist / contractor'],
 ['Ultrasonic report','Ultrasonic thickness, Ultrasonics (airborne)','Measured','Thickness: Lower is worse vs minimum wall. Airborne: dB limits','Quarterly / annual','UT contractor / RCM'],
 ['Oil analysis','Lubrication / oil','Measured','Lab alarm limits per parameter (water, iron, silicon, viscosity...)','Monthly / per service','RCM Specialist'],
 ['DG checklist','DG checks','Measured','Enter only the condition parameters (exhaust and turbo temperatures, oil pressure, coolant temperature, crankcase pressure) with limits','Daily; enter weekly summary or any exceedance','DG operator'],
 ['LV/HV inspection checklist','Fleet inspection','Checklist','Severity rules per component','Per service / monthly','Fleet mechanic'],
 ['LV/HV prestart checklist','Prestart','Exceptions only','Enter only defects found (severity rules); completion rate tracked separately','Every shift','Operator / supervisor'],
])
put('Severity rules for checklists, visual and prestart items',bold=True,size=12,color=NAVY)
table(['Status','Definition','Required response','Examples'],[
 ['OK','Item within standard','None (record only if part of a route)','Guards fitted, no leaks'],
 ['Alert','Defect, but safe to operate','Repair within P2/P3 response time','Minor oil leak, silica gel discoloured, worn seat, reverse alarm faulty'],
 ['Danger','Unsafe or failure imminent','Stop / isolate / tag out; repair before use (P1)','Brakes or steering ineffective, seatbelt or fire suppression faulty, major leak, exposed live parts'],
])
put('Asset numbering convention (site tag)',bold=True,size=12,color=NAVY)
table(['Prefix','Area','Example'],[['AP-','Acid plant','AP-P-101 (pump), AP-TK-01 (tank)'],['CK-','Calciner','CK-1080'],['SS-','Substation','SS-TX-01, SS-MCC-01'],['PP-','Power plant','PP-GEN-04'],['LV- / HV-','Mobile fleet','LV-01 light vehicle, HV-02 compactor']])
para('The Pronto asset no. remains the unique key; the site tag is the name people use in the field. Equipment ID is used as the asset key.')
row[0]+=1
put('Repeat readings rule',bold=True,size=12,color=NAVY)
para('Each new reading is a new row. If a reading repeats an issue already raised in an earlier row, set its Record stage to "No action": it updates the trend and the asset status, but does not create a second open recommendation. An asset returns to OK when its recommendation is Closed and the latest reading is within limits.')
row[0]+=1
put('Dashboard sheets',bold=True,size=12,color=NAVY)
table(['Sheet','Feeds'],[['Events','Machine history, MTBF/MTTR, availability, bad actors (Pronto WO export)'],['Schedule','Inspection schedule compliance'],['CM_Value / Programme_Cost','CM return on investment'],['Decisions','Management decisions required'],['Prestart','Fleet prestart compliance']])
put('Settings (Lists sheet)',bold=True,size=12,color=NAVY)
para('ACI weights, criticality class limits, priority bands and response times (P1 7 days, P2 30, P3 60), severity factors (Danger 1.0, Alert 0.6) and the statutory alert window (60 days) are editable on the Lists sheet. Get them approved by the Engineering Manager before go-live so the numbers are not disputed.')

# ---------------- KPI TARGETS (Lists) ----------------
ptab(34,['KPI target','Value'],[['Asset Health Index',85],['Inspection schedule compliance %',95],['Recommendations closed on time %',85],['Overdue recommendations (max)',2],['CM return on investment (x : 1)',4],['Critical asset CM coverage %',90],['Prestart compliance %',95]],'Targets shown on the Management page. Approve with the Engineering Manager.')

ptab(44,['Inspection interval','Days'],[['A',30],['B',60],['C',90]],'Maximum days between inspections by criticality class. Older = Unknown.')
# ---------------- EXTRA SHEETS ----------------
def simple_sheet(name,title,sub,cols,rows,datecols=(),money=(),tab='9CB6C6'):
    ws=wb.create_sheet(name)
    header(ws,[(c,w,'in') for c,w in cols],title,sub)
    for i,row in enumerate(rows):
        for j,v in enumerate(row,1):
            if v=='' or v is None: continue
            c=ws.cell(5+i,j,dt.datetime.strptime(v,'%Y-%m-%d') if (j in datecols and isinstance(v,str)) else v)
            c.font=SAMPLE; c.border=B
            if j in datecols: c.number_format='dd-mmm-yy' if 'Month' not in cols[j-1][0] else 'mmm-yy'
            if j in money: c.number_format='$#,##0'
    ws.freeze_panes='A5'; ws.sheet_properties.tabColor=tab
    return ws
simple_sheet('Events','AHIM Event History (from Pronto)','One row per failure, repair, PM or shutdown from Pronto work orders. Downtime only on Failure rows. Used for MTBF, MTTR, availability and bad actors.',
 [('Date',11),('Pronto asset no.',14),('Event type',11),('Description',44),('Downtime (h)',10),('Pronto WO no.',12),('RCA reference',13)],
 ev,datecols=(1,))
simple_sheet('Schedule','AHIM Inspection Schedule Compliance','One row per month and technique: inspections planned vs completed (from the CM schedule / Pronto PM routes).',
 [('Month',10),('Technique',22),('Planned',9),('Completed',10)],sch,datecols=(1,))
simple_sheet('CM_Value','AHIM Condition Monitoring Value (cost avoidance)','Net avoided = avoided failure cost (repair + collateral damage + lost production) minus planned repair cost. Count only when approved.',
 [('Date',11),('Pronto asset no.',14),('Technique',20),('Detection / failure avoided',46),('Avoided failure cost (USD)',14),('Planned repair cost (USD)',14),('Approved by',18)],
 OLD.get('CM_Value',[]),datecols=(1,),money=(5,6))
simple_sheet('Programme_Cost','AHIM CM Programme Cost','Monthly cost of the condition monitoring and inspection programme (USD).',
 [('Month',10),('Labour',11),('Lab & consumables',14),('Contractors',12),('Equipment',11)],OLD.get('Programme_Cost',[]),datecols=(1,),money=(2,3,4,5))
simple_sheet('Decisions','AHIM Management Decisions','Decisions only management can make (stops, budget, statutory bookings). Status: Open, Approved, Rejected or Closed.',
 [('Date raised',11),('Decision required',42),('Reason',60),('Owner',26),('Status',10)],OLD.get('Decisions',[]),datecols=(1,))
simple_sheet('Prestart','AHIM Prestart Compliance','Per month and vehicle/machine: shifts operated vs prestart checklists completed. Defects found go to Records.',
 [('Month',10),('Pronto asset no.',14),('Shifts operated',12),('Prestarts completed',14)],OLD.get('Prestart',[]),datecols=(1,))


_gen = {'Work_Orders': PRONTO_WO}
if not AS.old_sheet_rows(OUT, 'KPI_Tree', 8): _gen['KPI_Tree'] = [list(x) for x in AI.KPI_TREE]
AS.write_input_sheets(wb, _gen, OUT, INPUT_FILES, log=lambda n, m: rv('Inputs', n, m))
REVIEW.sort(key=lambda x: {'Summary':0,'Method':1,'Not in source':2}.get(x[0],3))
RVW=wb.create_sheet('Import_Review')
RVW['A1']='AHIM Import Review'; RVW['A1'].font=font(bold=True,size=14,color=NAVY)
RVW['A2']='What was converted from the CM Monthly Report workbook, the assumptions made, and items to check. Regenerated on every import.'; RVW['A2'].font=font(size=9,color=GREY)
for j,(h,w) in enumerate([('Category',30),('Item',34),('Detail',120)],1):
    c=RVW.cell(4,j,h); c.font=font(bold=True,color='FFFFFF'); c.fill=fill(NAVY); RVW.column_dimensions[L(j)].width=w
for i,(c1,c2,c3) in enumerate(REVIEW,5):
    for j,v in enumerate((c1,c2,c3),1):
        c=RVW.cell(i,j,v); c.font=font(bold=(c1 in ('Summary','Method'))); c.alignment=Alignment(wrap_text=True,vertical='top'); c.border=B
RVW.freeze_panes='A5'; RVW.sheet_properties.tabColor='B57C00'
simple_sheet('Commentary','AHIM Monthly Analyst Commentary','One row per month: three short lines for management. Shown at the top of the Management page.',[('Month',10),('What changed',55),('Why',55),('What we are doing',55),('Author',22)],_cm,datecols=(1,),tab='2E8A57')
for _c in 'BCD':
    for _r in range(5,40): wb['Commentary'][f'{_c}{_r}'].alignment=Alignment(wrap_text=True,vertical='top')
simple_sheet('WO_Summary','AHIM Pronto Work-Order Summary','Per month and work type, from the Pronto WO export: drives the SMRP maintenance performance metrics. Previous months are kept on re-import.',[('Month',10),('Work type',26),('Total',9),('Complete',10),('In progress',11),('Not started',11)],WOSUM,datecols=(1,))
AS.add_register_validation(AR, 41, RN)
AS.add_record_validation(RC, 32, RR)
AS.write_reference(wb, OUT if 'OUT' in globals() else None)
for ws in (AR,RC): ws.sheet_properties.tabColor=NAVY
G.sheet_properties.tabColor='2E8A57'; LS.sheet_properties.tabColor='9CB6C6'
wb.save(OUT)
print('Wrote',OUT)
print('ok',NA,NR)
